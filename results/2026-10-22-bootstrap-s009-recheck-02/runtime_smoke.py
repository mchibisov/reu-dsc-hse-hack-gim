"""Реальный CLI/SQL/сервис; только внешняя граница VK синтетическая."""
import asyncio
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import time
from contextlib import AsyncExitStack
from uuid import uuid4

import psycopg
from psycopg import sql
from psycopg.conninfo import conninfo_to_dict, make_conninfo
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb

SOURCE=Path(sys.argv[1]).resolve()
ARTIFACTS=Path(sys.argv[2]).resolve() if len(sys.argv)>2 else Path(__file__).resolve().parent
sys.path.insert(0,str(SOURCE))
sys.path.insert(0,str(SOURCE/'database/src'))
from gim.shared_config import SharedSettings
from gim.shared_gateway import enqueue_events
from gim.shared_main import check_database
from gim.shared_service import SharedService
from gim.shared_storage import LocalPrivateStorage
from gim_db.outbox import SQLQueue, run_once
from gim_db.repository import Repository, create_pool

ANSWERS={
    'consent':'Согласен','full_name':'Тестовый Кандидат','email':'synthetic@example.invalid',
    'vk_url':'https://vk.com/id123456','birth_date':'01.01.2000','current_status':'Студент',
    'university':'Тестовый университет','faculty':'Исторический, музеология','course':'2',
    'interests':'История','experience':'Помогал на тестовом мероприятии','languages':'Английский',
    'motivation':'Хочу помогать музею','expectations':'Получить опыт','available_time':'4 часа в неделю',
    'hobbies':'Фотография','directions':'Медиа','discovery':'Тест'}


class SyntheticVK:
    def __init__(self):
        self.messages=[]
        self.uploads=[]
        self.cmids={}

    def next_cmid(self, actor):
        self.cmids[actor]=self.cmids.get(actor,0)+1
        return self.cmids[actor]

    async def send(self,actor,text,random_id,**kwargs):
        self.messages.append((actor,text,kwargs))
        return {'conversation_message_id':self.next_cmid(actor)}

    async def upload(self,actor,name,data):
        self.uploads.append(data)
        return 'doc1_1'


async def scenario(settings,actor,admin):
    vk=SyntheticVK()
    async with AsyncExitStack() as stack:
        pools=[await stack.enter_async_context(create_pool(dsn,max_size=2)) for dsn in
               (settings.database_bot_url,settings.database_worker_url,settings.database_staff_url)]
        assert (await check_database(settings,pools))['base_questions']==18
        repositories=[Repository(pool) for pool in pools]
        bot,worker,staff=repositories
        storage=LocalPrivateStorage(SOURCE.parent/'s009-smoke-files')
        stack.push_async_callback(storage.close)
        service=SharedService(*repositories,vk,settings,storage)
        queue=SQLQueue(pools[1],settings.group_id)
        topics=['vk.incoming','question.send','vk.send','staff.notify','file.export']

        async def receive(actor,text):
            event={'type':'message_new','group_id':settings.group_id,'event_id':str(uuid4()),
                   'object':{'message':{'from_id':actor,'peer_id':actor,'text':text,
                                        'conversation_message_id':vk.next_cmid(actor)}}}
            async with pools[0].connection() as c,c.transaction():
                await enqueue_events(c,[event],settings.group_id,'')
            for _ in range(50):
                if not await run_once({'repository':worker,'handlers':service.handlers},queue,limit=10,topics=topics):
                    break
            else:
                raise AssertionError('Queue did not settle')
            async with pools[1].connection() as c:
                errors=await (await c.execute("""SELECT topic,last_error_code FROM gim.outbox_events
                    WHERE payload->>'group_id'=%s AND last_error_code IS NOT NULL""",(str(settings.group_id),))).fetchall()
                assert not errors,errors

        await receive(actor,'/start')
        app=await bot.start_application(settings.recruitment_id,actor)
        for slot,answer in ANSWERS.items():
            context=await bot.dialog_context(app['id'])
            question=next(q for q in context['questions'] if q['id']==context['application']['current_question_id'])
            assert question['target_slot']==slot,(slot,question['target_slot'])
            await receive(actor,answer)
        context=await bot.dialog_context(app['id'])
        assert context['application']['status']=='submitted'
        assert len(context['answers'])==18
        assert all(q['target_slot']!='university_other' for q in context['questions'])
        prompts=[text for peer,text,kw in vk.messages if peer==actor and kw.get('want_cmid')]
        assert len(prompts)==18
        assert any(text.endswith('Факультет, направление') for text in prompts)
        await receive(admin,'/export')
        assert len(vk.uploads)==1
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(vk.uploads[0]),read_only=True)
        assert set(workbook.sheetnames)=={'Анкеты','Ответы','Собеседование','Сводка','Оценка'}
        result={'submitted':True,'answers':18,'delivered_prompts':18,'deleted_question_absent':True,
                'xlsx_sheets':workbook.sheetnames,'xlsx_bytes':len(vk.uploads[0]),
                'external_vk_llm_called':False,'analysis':'not executed; no external key in acceptance test'}
        workbook.close()
        return result


def main():
    dsn=os.environ['TEST_DATABASE_URL']
    params=conninfo_to_dict(dsn)
    assert params['host']=='127.0.0.1' and params['dbname']=='gim_s009_second'
    suffix=uuid4().hex[:8]
    group=uuid4().int%1000000000+1
    actor=group+10000000000
    admin=actor+1
    env={**os.environ,'VK_GROUP_ID':str(group),'VK_TOKEN':'synthetic-not-a-real-token',
         'LLM_API_KEY':'','ANALYSIS_RUBRIC_PATH':'','MAX_FOLLOWUPS':'0','EXPECTATIONS_CHECK':'false',
         'STORAGE_DIR':str(SOURCE.parent/'s009-smoke-files'),'PYTHONIOENCODING':'utf-8'}
    logins=[]
    process=None
    with psycopg.connect(dsn,autocommit=True,row_factory=dict_row) as c:
        try:
            for role in ('bot','worker','staff'):
                login='s009_'+role+'_'+suffix
                c.execute(sql.SQL('CREATE ROLE {} LOGIN INHERIT IN ROLE {}').format(sql.Identifier(login),sql.Identifier('gim_'+role)))
                logins.append(login)
                env['DATABASE_'+role.upper()+'_URL']=make_conninfo(**{**params,'user':login})
            recruitment=c.execute("""INSERT INTO gim.recruitments
                (group_id,title,questionnaire_version,rubric_version,rubric,followup_policy)
                VALUES(%s,'Synthetic clean start','s009','s009',%s,%s) RETURNING id""",
                (group,Jsonb({'criteria':[],'approval':'synthetic-only'}),Jsonb({'max_followups':0,'expectations_check':False}))).fetchone()['id']
            for position,slot in enumerate(ANSWERS,1):
                c.execute("""INSERT INTO gim.question_templates(recruitment_id,code,kind,text,position,required,target_slot)
                    VALUES(%s,%s,'base',%s,%s,true,%s)""",(recruitment,slot,'Факультет, направление' if slot=='faculty' else slot,position,slot))
            c.execute("UPDATE gim.recruitments SET status='active',opened_at=now() WHERE id=%s",(recruitment,))
            user=c.execute('INSERT INTO gim.users(vk_id) VALUES(%s) RETURNING id',(admin,)).fetchone()['id']
            c.execute("INSERT INTO gim.staff_members(user_id,group_id,role) VALUES(%s,%s,'admin')",(user,group))
            env['RECRUITMENT_ID']=str(recruitment)
            check=subprocess.run([sys.executable,'-m','gim.shared_main','check-db'],cwd=SOURCE,env=env,capture_output=True,text=True,timeout=25)
            (ARTIFACTS/'check-db.log').write_text(check.stdout+check.stderr,encoding='utf-8')
            assert check.returncode==0,check.returncode
            assert json.loads(check.stdout)['base_questions']==18
            with (ARTIFACTS/'worker.log').open('w',encoding='utf-8') as log:
                process=subprocess.Popen([sys.executable,'-m','gim.shared_main','worker'],cwd=SOURCE,env=env,
                    stdout=log,stderr=subprocess.STDOUT,creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
                for _ in range(100):
                    if process.poll() is not None:raise AssertionError('Worker exited before startup')
                    if 'shared_runtime_start' in (ARTIFACTS/'worker.log').read_text(encoding='utf-8'):break
                    time.sleep(.1)
                else:raise AssertionError('Worker startup timeout')
                time.sleep(1)
                assert process.poll() is None
                process.terminate()
                process.wait(timeout=10)
                process=None
            os.environ.update(env)
            result=asyncio.run(scenario(SharedSettings.load(),actor,admin))
            result.update(cli_check_db=True,worker_process_started=True,source=str(SOURCE))
            print(json.dumps(result,ensure_ascii=False))
        finally:
            if process is not None:
                process.terminate()
                process.wait(timeout=10)
            for login in logins:
                c.execute(sql.SQL('DROP ROLE {}').format(sql.Identifier(login)))


if __name__=='__main__':
    if os.name=='nt':asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    main()
