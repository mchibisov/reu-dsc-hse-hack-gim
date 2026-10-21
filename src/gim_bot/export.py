from collections import Counter
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .questionnaire import QUESTIONS

STATUS = {'queued': 'Ожидает обработки', 'processing': 'Обрабатывается',
          'ready': 'Готова', 'no_analytics': 'Без аналитики', 'pending': 'Не завершена'}


def build_excel(applications, total_started=None) -> bytes:
    apps = sorted(applications, key=lambda a: (a.score is None, -(a.score or 0), a.id))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Анкеты'
    sheet.append(['ID', 'Получена', 'Балл 0–3', 'Аналитика', 'Рубрика',
                  *[q.text for q in QUESTIONS], 'Уточнения и ответы', 'Резюме',
                  'Рекомендуемое направление', 'Обоснование направления', 'Критерии и цитаты',
                  'Вопросы для собеседования', 'Неполнота оценки'])
    for app in apps:
        analysis = app.analysis or {}
        criteria = '\n'.join(f"{c['id']}: {c['level']} — {c['explanation']} [{c['field']}: {c['quote']}]"
                             for c in analysis.get('criteria', []))
        followups = '\n\n'.join(f"{q['question']}\n{q.get('answer', '')}" for q in app.followups)
        sheet.append([app.id, app.completed_at.isoformat() if app.completed_at else '', app.score,
            STATUS.get(app.analysis_status, app.analysis_status), app.rubric_version or '',
            *[app.answers.get(q.key, '') for q in QUESTIONS], followups,
            analysis.get('summary', ''), analysis.get('recommended_direction', ''),
            analysis.get('rationale', ''), criteria, '\n'.join(analysis.get('interview_questions', [])),
            'Требуется проверка' if analysis.get('requires_review') or app.score is None else ''])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 64
    for col in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 38 if col > 5 else 20
    summary = workbook.create_sheet('Сводка')
    summary.append(['Показатель', 'Значение'])
    summary.append(['Начато анкет', total_started if total_started is not None else len(apps)])
    summary.append(['Завершено', len(apps)])
    summary.append(['Без балла / в обработке', sum(a.score is None for a in apps)])
    counts = Counter(a.score for a in apps if a.score is not None)
    for score in range(4):
        summary.append([f'Балл {score}', counts[score]])
    for direction, count in Counter((a.analysis or {}).get('recommended_direction', 'Не определено') for a in apps).items():
        summary.append([direction, count])
    summary.append(['Примечание', 'Решение о приёме принимает сотрудник. Рубрика предварительная.'])
    summary.column_dimensions['A'].width = 40
    summary.column_dimensions['B'].width = 75
    for worksheet in workbook:
        for row in worksheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Даже '=...' остаётся текстом, без исполняемых формул.
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='264A65')
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
