"""XLSX строго из сохранённого снимка; продолжения не теряют исходный текст."""
import io
import json
import re
import zipfile
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font

SHEETS=('Анкеты','Ответы','Собеседование','Сводка')


def _text(value):
    if value is None:
        return ''
    if isinstance(value,(dict,list)):
        return json.dumps(value,ensure_ascii=False,default=str,sort_keys=True)
    return str(value)


def snapshot_xlsx(snapshot):
    wb=Workbook()
    wb.remove(wb.active)
    wb.properties.created=wb.properties.modified=datetime(2000,1,1)
    for name in SHEETS:
        sheet=wb.create_sheet(name)
        data=snapshot.get(name,[] if name!='Сводка' else {})
        if isinstance(data,dict):
            data=[dict(metric=k,value=v) for k,v in data.items()]
        prepared=[]
        for record in data:
            record=dict(record)
            encoded={}
            for key,value in list(record.items()):
                text=_text(value)
                if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',text):
                    record[key]=json.dumps(text,ensure_ascii=False)
                    encoded[key]='json_string'
            if encoded:
                record['field_encodings']=encoded
            prepared.append(record)
        data=prepared
        columns=list(dict.fromkeys(k for row in data for k in row)) or ['Нет данных']
        columns += ['continuation_part','continuation_total']
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font=Font(bold=True)
        for row in data:
            # 15000 code points also fit Excel's UTF-16 limit for astral emoji.
            chunks={k:[_text(row.get(k))[i:i+15000] for i in range(0,len(_text(row.get(k))),15000)] or [''] for k in columns[:-2]}
            total=max(map(len,chunks.values()))
            for n in range(total):
                values=[]
                for k in columns[:-2]:
                    # Carry identifiers on continuation rows; split each other field losslessly.
                    v=chunks[k]
                    values.append(v[n] if n<len(v) else (v[0] if k.endswith('_id') else ''))
                sheet.append(values+[str(n+1),str(total)])
                for cell in sheet[sheet.max_row]:
                    cell.data_type='s'
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
    data=io.BytesIO()
    wb.save(data)
    # Fixed ZIP timestamps make crash retry byte-identical for immutable storage.
    output=io.BytesIO()
    with zipfile.ZipFile(data) as source,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            item.date_time=(2000,1,1,0,0,0)
            content=source.read(item.filename)
            if item.filename=='docProps/core.xml':
                # openpyxl updates modified during save regardless of properties.
                content=re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)',
                               rb'\g<1>2000-01-01T00:00:00Z\g<2>',content)
            target.writestr(item,content)
    return output.getvalue()
