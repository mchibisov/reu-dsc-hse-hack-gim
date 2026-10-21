"""Idempotent import of museum-labelled historical forms; never sends VK messages."""
import argparse
import asyncio
import hashlib
import json
import math
import uuid
from datetime import date, datetime
from pathlib import Path

from psycopg.types.json import Jsonb

from gim.config import Settings, load_dotenv
from gim.db import Database
from gim.dialog import DIRECTIONS

FIELDS = {
    'ФИО': 'name', 'Дата рождения': 'birth_date', 'ВУЗ': 'education', 'Факультет': 'faculty',
    'Курс': 'course', 'Кафедра': 'department', 'Сфера научных интересов': 'interests',
    'Опыт работы': 'experience', 'Владение иностранными языками': 'languages',
    'Интересующие виды волонтёрства': 'directions', 'Чего Вы ожидаете': 'motivation',
    'Ваше хобби': 'hobbies', 'В чем Ваша суперсила': 'strengths', 'Откуда Вы узнали': 'source',
}
OUTCOMES = {'ЯВКА': 'attendance', 'Взяли ли': 'accepted', 'Длительность': 'years',
            'Еще состоит': 'still_active', 'Активность в часах': 'hours'}


def normalize_record(headers, row):
    answers, outcomes, score, name, directions = [], {}, None, None, []
    for header, value in zip(headers, row):
        if header is None or value is None or (isinstance(value, str) and not value.strip()):
            continue
        header = ' '.join(str(header).split())
        if isinstance(value, (datetime, date)):
            value = value.isoformat()
        if header.startswith('Сколько баллов'):
            score = float(value)
            if not math.isfinite(score) or not 0 <= score <= 3:
                raise ValueError('Historical score outside 0–3')
            continue
        outcome = next((code for prefix, code in OUTCOMES.items() if header.startswith(prefix)), None)
        if outcome:
            outcomes[outcome] = value
            continue
        code = next((code for prefix, code in FIELDS.items() if header.startswith(prefix)), None)
        if not code:
            continue
        if code == 'name':
            name = str(value).strip()
        text = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()
        answers.append({'code': code, 'question': header, 'text': text})
        if code == 'directions':
            directions = [d for d, (_, terms) in DIRECTIONS.items() if any(not term.isdigit() and term in text.casefold() for term in terms)]
    if not answers:
        return None
    return {'answers': answers, 'outcomes': outcomes, 'score': score, 'name': name, 'directions': directions}


def read_history(path):
    from openpyxl import load_workbook
    path = Path(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, data_only=True, read_only=True)
    result = []
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            for number, row in enumerate(rows, start=2):
                record = normalize_record(headers, row)
                if record:
                    result.append({**record, 'source': {'file': path.name, 'sha256': digest,
                                                       'sheet': sheet.title, 'row': number}})
    finally:
        workbook.close()
    return result


async def import_records(db, records, group_id, campaign):
    inserted = 0
    async with db.pool.connection() as c:
        for record in records:
            source = record['source']
            identity = f"gim:{group_id}:{campaign}:{source['sha256']}:{source['sheet']}:{source['row']}"
            app_id = uuid.uuid5(uuid.NAMESPACE_URL, identity)
            user_id = uuid.uuid5(uuid.NAMESPACE_OID, str(app_id))
            await c.execute('INSERT INTO users(id,name) VALUES(%s,%s) ON CONFLICT(id) DO NOTHING', (user_id, record['name']))
            app = await (await c.execute('''INSERT INTO applications(id,user_id,group_id,campaign,status,revision,state,profile,selected_directions)
                VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(id) DO NOTHING RETURNING id''',
                (app_id, user_id, group_id, campaign, 'evaluated' if record['score'] is not None else 'submitted', len(record['answers']),
                 Jsonb({'step': 'submitted', 'historical': True}), Jsonb({'source': source, 'historical_outcomes': record['outcomes']}), record['directions']))).fetchone()
            if not app:
                continue
            for revision, answer in enumerate(record['answers'], start=1):
                event_id = f'history:{app_id}:{revision}'
                await c.execute('INSERT INTO inbox_events(group_id,event_id,application_id) VALUES(%s,%s,%s)', (group_id, event_id, app_id))
                await c.execute('''INSERT INTO answers(application_id,question_code,question_kind,question_text,answer_text,revision,group_id,event_id)
                    VALUES(%s,%s,'base',%s,%s,%s,%s,%s)''', (app_id, answer['code'], answer['question'], answer['text'], revision, group_id, event_id))
            if record['score'] is not None:
                await c.execute('''INSERT INTO evaluations(application_id,score,source,answer_revision,rubric_version,rubric_snapshot,summary)
                    VALUES(%s,%s,'human',%s,'historical-museum',%s,%s)''',
                    (app_id, record['score'], len(record['answers']), Jsonb({'source': source, 'scale': '0: not invited; 1: weak; 2: medium; 3: good'}), 'Историческая экспертная оценка музея'))
            await c.execute("INSERT INTO audit_log(application_id,action,details) VALUES(%s,'historical_import',%s)", (app_id, Jsonb(source)))
            inserted += 1
    return inserted


def main():
    load_dotenv()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('file', type=Path)
    parser.add_argument('--campaign', required=True)
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()
    records = read_history(args.file)
    counts = {str(score): sum(r['score'] == score for r in records) for score in range(4)}
    print(json.dumps({'records': len(records), 'scores': counts}, ensure_ascii=False))
    if not args.dry_run:
        async def run():
            settings = Settings.load()
            db = Database(settings.database_url)
            await db.open()
            try:
                print('inserted:', await import_records(db, records, settings.group_id, args.campaign))
            finally:
                await db.close()
        asyncio.run(run())


if __name__ == '__main__':
    main()
