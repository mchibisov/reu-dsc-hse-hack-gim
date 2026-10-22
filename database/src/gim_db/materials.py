"""Private material preparation: allowlisted pre-application fields only.

Outputs remain private. Expert label is kept separately from model features.
Never import outcome columns (hours, admission, tenure, still active).
"""
import hashlib
import json
import re
from datetime import date,datetime
from pathlib import Path

# Header fragments distinguish fields despite whitespace and line break differences.
FEATURE_COLUMNS={
 'Дата рождения':'birth_date_raw','ВУЗ':'university','Факультет':'faculty','Курс':'course',
 'Кафедра':'department','Сфера научных интересов':'interests','Опыт работы':'experience',
 'Владение иностранными языками':'languages','Интересующие виды волонт':'directions',
 'Чего Вы ожидаете':'expectations','Ваше хобби':'hobbies','В чем Ваша суперсила':'strengths','Откуда Вы узнали':'discovery',
}


def prepare_history(paths):
    from openpyxl import load_workbook
    records=[]; counts={}
    for path in paths:
        count=0
        for sheet in load_workbook(path,data_only=True,read_only=True).worksheets:
            rows=iter(sheet.values); headers=next(rows)
            mapping={}; score_col=None
            for i,h in enumerate(headers):
                normalized=' '.join(str(h or '').split())
                if normalized.startswith('Сколько баллов'):score_col=i
                for prefix,code in FEATURE_COLUMNS.items():
                    if normalized.startswith(prefix):mapping[i]=code
            if score_col is None:raise ValueError('Expert score header missing')
            for number,row in enumerate(rows,start=2):
                if not any(v is not None for v in row):continue
                features={code:serialize(row[i]) for i,code in mapping.items() if i<len(row)}
                score=row[score_col]
                if score is not None:
                    if isinstance(score,bool) or str(score).strip() not in {'0','1','2','3','0.0','1.0','2.0','3.0'}:
                        raise ValueError(f'Invalid expert score at sheet row {number}')
                    score=int(float(score))
                identity=f'{Path(path).name}:{sheet.title}:{number}'
                records.append({'id':hashlib.sha256(identity.encode()).hexdigest()[:24],
                    'cohort':Path(path).stem,'source_sheet':sheet.title,'source_row':number,
                    'features':features,'expert_score':score})
                count+=1
        counts[Path(path).name]=count
    return records,counts


def serialize(v):
    # No invented day/month for partial historical dates.
    if isinstance(v,(datetime,date)):return v.isoformat()
    if v is None:return None
    return str(v)


def write_private_jsonl(path,records):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    import os
    fd=os.open(path,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
    with os.fdopen(fd,'w',encoding='utf-8') as f:
        for r in records:f.write(json.dumps(r,ensure_ascii=False)+'\n')


def questionnaire_from_docx(path):
    """Extract this museum document; fail closed if its structure changes."""
    from docx import Document
    paragraphs=[p.text.strip() for p in Document(path).paragraphs if p.text.strip()]
    starts=['Даю согласие','ФИО','Электронная почта','Ссылка на страницу','Дата рождения','Ваш нынешний статус',
      'ВУЗ','ВУЗ (Другое)','Факультет','Курс','Сфера научных интересов','Опыт работы','Владение иностранными языками',
      'Почему Вы выбрали','Чего Вы ожидаете','Сколько времени','Ваши хобби','В работу какого направления','Откуда Вы узнали']
    codes=['consent','full_name','email','vk_url','birth_date','current_status','university','university_other',
      'faculty','course','interests','experience','languages','motivation','expectations','available_time','hobbies','directions','discovery']
    begin=next(i for i,p in enumerate(paragraphs) if p.startswith('Даю согласие'))
    end=next(i for i in range(begin,len(paragraphs)) if 'Критерии оценки' in paragraphs[i])
    bank=paragraphs[begin:end]
    result=[];cursor=0
    for position,(prefix,code) in enumerate(zip(starts,codes),1):
        found=next((i for i in range(cursor,len(bank)) if bank[i].startswith(prefix)),None)
        if found is None:raise ValueError(f'Question slot missing: {code}')
        result.append({'code':code,'kind':'base','text':bank[found],'position':position,'required':True,
            'target_slot':code,'validation':{'allow_not_applicable':code in {'university','university_other','faculty','course','languages','experience'}},'directions':[]})
        cursor=found+1
    # Include option paragraphs between consecutive question headings.
    offsets=[]; cursor=0
    for prefix in starts:
        found=next(i for i in range(cursor,len(bank)) if bank[i].startswith(prefix))
        offsets.append(found); cursor=found+1
    for i,q in enumerate(result):
        stop=offsets[i+1] if i+1<len(offsets) else len(bank)
        q['text']='\n'.join(bank[offsets[i]:stop])
    start=next(i for i,p in enumerate(paragraphs) if 'Уточняющие вопросы' in p)
    # The provided followup bank has question marks, unlike its category headings.
    followups=[p for p in paragraphs[start+1:] if p.endswith('?')]
    slots=[('collections_interest',['collections']),('alternatives',['collections']),('project_formats',['project']),('visitor_groups',['organizational']),('media_activity',['media']),('schedule',[]),('other_commitments',[]),('teamwork',[]),('qualities',[]),('expectation_goals',[]),('project_ideas',[])]
    for i,text in enumerate(followups):
        result.append({'code':f'followup_{i+1}','kind':'followup','text':text,'position':len(starts)+i+1,
           'required':False,'target_slot':slots[i][0],'directions':slots[i][1],'validation':{}})
    if len(followups)!=11:raise ValueError('Followup bank changed; review parser')
    return result


def direction_text_from_pdf(path):
    from pypdf import PdfReader
    # Preserve private original page excerpts for team mapping; do not invent requirements.
    return [{'page':i,'text':p.extract_text()} for i,p in enumerate(PdfReader(path).pages,1)]
