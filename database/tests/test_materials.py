from pathlib import Path
import pytest
from gim_db.materials import FEATURE_COLUMNS,prepare_history,questionnaire_from_docx


def test_private_jsonl_is_utf8_and_preserves_unicode(tmp_path):
    from gim_db.materials import write_private_jsonl
    import json
    path = tmp_path / 'history.jsonl'
    records = [{'id': 'synthetic', 'features': {'experience': 'Музей 🏛'}, 'expert_score': 2}]
    write_private_jsonl(path, records)
    assert [json.loads(line) for line in path.read_text(encoding='utf-8').splitlines()] == records


@pytest.mark.parametrize('include_old_question', [True, False])
def test_utf8_bom_recruitment_import_on_empty_database(tmp_path, include_old_question):
    import json
    import os
    import subprocess
    import sys
    import uuid
    import pytest
    import psycopg
    from docx import Document
    from pypdf import PdfWriter
    dsn = os.environ.get('TEST_DATABASE_URL')
    if not dsn:
        pytest.skip('disposable PostgreSQL required')
    root = Path(__file__).resolve().parents[2]
    document = Document()
    headings = ['Даю согласие', 'ФИО', 'Электронная почта', 'Ссылка на страницу',
                'Дата рождения', 'Ваш нынешний статус', 'ВУЗ', 'ВУЗ (Другое)',
                'Факультет', 'Курс', 'Сфера научных интересов', 'Опыт работы',
                'Владение иностранными языками', 'Почему Вы выбрали', 'Чего Вы ожидаете',
                'Сколько времени', 'Ваши хобби', 'В работу какого направления', 'Откуда Вы узнали']
    if not include_old_question:
        headings.remove('ВУЗ (Другое)')
    for heading in headings:
        document.add_paragraph(heading)
    document.add_paragraph('Критерии оценки')
    document.add_paragraph('Уточняющие вопросы')
    for number in range(11):
        document.add_paragraph(f'Синтетическое уточнение {number + 1}?')
    document.save(tmp_path / 'questionnaire.docx')
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(tmp_path / 'directions.pdf')
    rubric = json.loads((root / 'config/ranking-50-35-15.json').read_text(encoding='utf-8'))
    config = {'group_id': uuid.uuid4().int % 1000000000 + 1, 'title': 'Тестовый музей 🏛',
              'questionnaire_version': 'synthetic', 'rubric_version': rubric['version'],
              'rubric': rubric['rubric'], 'followup_policy': {'max_followups': 3},
              'direction_info': {key: 'Синтетическое описание' for key in
                                 ('organizational', 'project', 'collections', 'inclusive', 'media')}}
    config_path = tmp_path / 'recruitment.json'
    config_path.write_text(json.dumps(config, ensure_ascii=False), encoding='utf-8-sig')
    result = subprocess.run([sys.executable, 'database/scripts/load_materials.py',
                             '--docx', str(tmp_path / 'questionnaire.docx'),
                             '--directions-pdf', str(tmp_path / 'directions.pdf'),
                             '--config', str(config_path)], cwd=root,
                            env={**os.environ, 'DATABASE_ADMIN_URL': dsn},
                            capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
    imported = json.loads(result.stdout)
    assert imported['status'] == 'draft' and imported['templates'] == 29
    with psycopg.connect(dsn) as connection:
        row = connection.execute('SELECT title,rubric,direction_info FROM gim.recruitments WHERE id=%s',
                                 (imported['recruitment_id'],)).fetchone()
    assert row[0] == config['title'] and row[1] == config['rubric']
    assert row[2]['media'] == 'Синтетическое описание'

    # Finish the documented operator bootstrap with actual least-privilege LOGINs.
    import secrets
    credentials = {role: {'login': f'gim_{role}_{uuid.uuid4().hex[:12]}',
                          'password': secrets.token_urlsafe(24)}
                   for role in ('bot', 'worker', 'staff')}
    credentials_path = tmp_path / 'credentials.json'
    credentials_path.write_text(json.dumps(credentials), encoding='utf-8-sig')
    operator_env = {**os.environ, 'DATABASE_ADMIN_URL': dsn}
    try:
        for args in (['database/scripts/provision_runtime.py', str(credentials_path)],
                     ['database/scripts/bootstrap_admin.py', '--group-id', str(config['group_id']),
                      '--verified-vk-id', str(config['group_id'] + 2000000000)]):
            setup = subprocess.run([sys.executable, *args], cwd=root, env=operator_env,
                                   capture_output=True, text=True, timeout=30)
            assert setup.returncode == 0, setup.stderr
        with psycopg.connect(dsn) as connection:
            connection.execute("UPDATE gim.recruitments SET status='active',opened_at=now() WHERE id=%s",
                               (imported['recruitment_id'],))
        from psycopg.conninfo import conninfo_to_dict, make_conninfo
        runtime_env = {**os.environ, 'VK_GROUP_ID': str(config['group_id']),
                       'RECRUITMENT_ID': imported['recruitment_id']}
        runtime_env.pop('DATABASE_ADMIN_URL', None)
        for role, credentials_entry in credentials.items():
            connection_info = {**conninfo_to_dict(dsn), 'user': credentials_entry['login'],
                               'password': credentials_entry['password']}
            runtime_env[f'DATABASE_{role.upper()}_URL'] = make_conninfo(**connection_info)
        checked = subprocess.run([sys.executable, '-m', 'gim.shared_main', 'check-db'],
                                 cwd=root, env=runtime_env, capture_output=True, text=True, timeout=30)
        assert checked.returncode == 0, checked.stderr
        assert json.loads(checked.stdout)['base_questions'] == 18
    finally:
        from psycopg import sql
        with psycopg.connect(dsn, autocommit=True) as connection:
            for entry in credentials.values():
                connection.execute(sql.SQL('DROP ROLE IF EXISTS {}').format(sql.Identifier(entry['login'])))


def test_prediction_metrics_read_utf8_bom_jsonl(tmp_path):
    import json
    import subprocess
    import sys
    root = Path(__file__).resolve().parents[2]
    truth = tmp_path / 'history.jsonl'
    predictions = tmp_path / 'predictions.jsonl'
    truth.write_text(json.dumps({'id': 'анкета 🏛', 'expert_score': 2}, ensure_ascii=False) + '\n',
                     encoding='utf-8-sig')
    predictions.write_text(json.dumps({'id': 'анкета 🏛', 'score': 2}, ensure_ascii=False) + '\n',
                           encoding='utf-8-sig')
    result = subprocess.run([sys.executable, 'database/scripts/evaluate_predictions.py',
                             str(truth), str(predictions)], cwd=root,
                            capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout) == {'predicted': 1, 'available_labels': 1,
                                       'coverage': 1, 'exact_match': 1, 'mae': 0}


def test_future_outcomes_are_never_features(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'synthetic.xlsx';w=Workbook();s=w.active
    s.append(['Сколько баллов дали','Дата рождения','Опыт работы','Активность в часах в год','Длительность пребывания в программе','Взяли ли в команду'])
    s.append([0,'01.2000','Снимал',999,42,'Да']);w.save(p)
    records,counts=prepare_history([p])
    assert records[0]['features']=={'birth_date_raw':'01.2000','experience':'Снимал'}
    assert records[0]['expert_score']==0 and sum(counts.values())==1
