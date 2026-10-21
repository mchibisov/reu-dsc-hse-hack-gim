"""Opt-in synthetic ingress → real Supabase/LLM → real curator VK/XLSX check.

Run only in an isolated test group/DB, or with its matching normal worker stopped. Never stop production for development. Candidate messages are captured locally;
staff messages/files go only to --staff-vk-id, which must already be authorized.
Creates two clearly labelled synthetic applications; never edits existing forms.
Invoke as python -m scripts.smoke_api --staff-vk-id YOUR_ID.
"""
import argparse
import asyncio
import hashlib
import io
import json
import os
import time
from contextlib import AsyncExitStack,suppress
from pathlib import Path
from uuid import uuid4

import httpx
from openpyxl import load_workbook

from gim.config import load_dotenv
from gim.shared_config import SharedSettings
from gim.shared_gateway import GatewayLedger
from gim.shared_main import check_database
from gim.shared_service import SharedService
from gim.shared_storage import LocalPrivateStorage
from gim.vk import VK
from gim_db.outbox import SQLQueue,run_once
from gim_db.repository import Repository,create_pool,one


COMMON={'consent':'Согласен','email':'synthetic@example.test','vk_url':'https://vk.com/synthetic',
    'birth_date':'01.01.2000','current_status':'Работаю неполный день, учёбу закончил',
    'university':'РГГУ','university_other':'не относится','faculty':'Исторический факультет',
    'course':'Бакалавриат окончен','languages':'Английский B2',
    'available_time':'Живу в Москве, могу приезжать в ГИМ по субботам. 16 часов в месяц, включая работу на месте. Других волонтёрских программ нет, график согласую с куратором.',
    'discovery':'От знакомого музейного волонтёра'}
PROFILES=[{**COMMON,'full_name':'ДЕМО — медиа (синтетическая анкета)',
    'interests':'История России XIX века и музейные медиапроекты',
    'experience':'Два года волонтёрил в краеведческом музее. Лично написал 20 постов о выставках и контент-план на месяц, фотографировал экспонаты с разрешения хранителя. Работал в редакционной команде.',
    'motivation':'Хочу помогать ГИМ рассказывать историю России через точные публикации о коллекциях.',
    'expectations':'Готовить рубрику об истории быта вместе с редактором, сверять факты с хранителями, получать правки. Интересуют посты и фото.',
    'hobbies':'Фотография и обработка фото в Lightroom','directions':'медиа'},
    {**COMMON,'full_name':'ДЕМО — фонды (синтетическая анкета)',
    'interests':'Археология, нумизматика, систематизация музейных коллекций',
    'experience':'На практике в архиве сверял описи, вносил номера в электронную таблицу и сканировал документы. Лично обработал 50 карточек; находил расхождения и обсуждал их с руководителем.',
    'motivation':'Хочу помогать ГИМ делать коллекции доступнее исследователям через оцифровку.',
    'expectations':'Интересны археологический отдел и книжный фонд, хочу освоить КАМИС и составление каталогов. Рассматриваю и помощь на мероприятиях, готов к обучению.',
    'hobbies':'Краеведение и чтение каталогов','directions':'фондовое и организационное'}]
CLARIFICATION='Мой личный вклад — сверка источников и аккуратное оформление результата. Согласовывал спорные данные с руководителем, проверял номера и исправлял ошибки. Готов работать в команде, осваивать новые задачи и регулярно приезжать в музей.'


class CheckedVK(VK):
    def __init__(self,settings,staff):
        super().__init__(settings)
        self.staff=staff
        self.synthetic=set()
        self.messages=[]

    async def send(self,peer_id,text,random_id,**kwargs):
        if peer_id in self.synthetic:
            self.messages.append({'peer_id':peer_id,'text':text})
            return random_id
        if peer_id!=self.staff:
            raise ValueError('Smoke must not deliver to third parties')
        return await super().send(peer_id,text,random_id,**kwargs)

    async def upload(self,peer_id,name,data):
        if peer_id!=self.staff:
            raise ValueError('Smoke must not upload for third parties')
        return await super().upload(peer_id,name,data)


async def smoke(staff_vk_id,output):
    settings=SharedSettings.load()
    async with AsyncExitStack() as stack:
        pools=[]
        for dsn in (settings.database_bot_url,settings.database_worker_url,settings.database_staff_url):
            pool=await stack.enter_async_context(create_pool(dsn))
            await pool.wait()
            pools.append(pool)
        await check_database(settings,pools)
        vk=CheckedVK(settings,staff_vk_id)
        stack.push_async_callback(vk.close)
        storage=LocalPrivateStorage(settings.storage_dir)
        stack.push_async_callback(storage.close)
        service=SharedService(*(Repository(p) for p in pools),vk,settings,storage)
        if not await service._staff(staff_vk_id):
            raise ValueError('Explicit staff authorization required')
        ledger=GatewayLedger(pools[0],settings.group_id,settings.excel_access_code)
        await ledger.acquire()  # excludes the regular gateway while verifying ingress
        stack.push_async_callback(ledger.release)
        async def heartbeat():
            while True:
                await asyncio.sleep(20)
                await ledger.execute('SELECT 1')
        lease_task=asyncio.create_task(heartbeat())
        async def stop_heartbeat():
            lease_task.cancel()
            with suppress(asyncio.CancelledError):
                await lease_task
        stack.push_async_callback(stop_heartbeat)
        queue=SQLQueue(pools[1],settings.group_id)
        ctx={'repository':service.worker_repo,'handlers':service.handlers}
        report={'ingress':'synthetic','candidate_vk_delivery':'captured; not a browser test',
                'staff_vk_delivery':'real','applications':[]}
        async with pools[0].connection() as c:
            fixtures=await (await c.execute('''SELECT u.vk_id FROM gim.applications a JOIN gim.users u ON u.id=a.user_id
                WHERE a.recruitment_id=%s AND u.full_name=ANY(%s::text[]) AND u.vk_id>=9000000000 AND u.vk_id<10000000000''',
                (settings.recruitment_id,[p['full_name'] for p in PROFILES]))).fetchall()
        vk.synthetic.update(r['vk_id'] for r in fixtures)

        async def drain():
            deadline=time.monotonic()+180
            while time.monotonic()<deadline:
                if await run_once(ctx,queue,limit=10):
                    continue
                async with pools[1].connection() as c:
                    waiting=await one(c,'''SELECT count(*) AS n FROM gim.outbox_events e
                        LEFT JOIN gim.applications a ON a.id=e.application_id
                        LEFT JOIN gim.recruitments r ON r.id=a.recruitment_id
                        LEFT JOIN gim.files f ON f.id=e.file_id
                        LEFT JOIN gim.recruitments fr ON fr.id=f.recruitment_id
                        WHERE COALESCE(r.group_id::text,fr.group_id::text,e.payload->>'group_id')=%s
                        AND e.completed_at IS NULL AND e.failed_at IS NULL''',(str(settings.group_id),))
                if not waiting['n']:
                    return
                # A retry delayed by backoff is still pending; do not feed the
                # next answer until the failed turn has actually recovered.
                await asyncio.sleep(0.5)
            raise RuntimeError('Queue did not drain')

        async def incoming(actor,text):
            event={'type':'message_new','group_id':settings.group_id,'event_id':'api-smoke:'+uuid4().hex,
                   'object':{'message':{'from_id':actor,'peer_id':actor,'text':text}}}
            # Do not alter the real Long Poll cursor with synthetic updates.
            ts=await ledger.cursor(settings.group_id)
            await ledger.ingest_batch([event],settings.group_id,ts or '0')
            await ledger.ingest_batch([event],settings.group_id,ts or '0')
            await drain()

        for values in PROFILES:
            async with pools[0].connection() as c:
                existing=await one(c,'''SELECT u.vk_id FROM gim.applications a JOIN gim.users u ON u.id=a.user_id
                    WHERE a.recruitment_id=%s AND u.full_name=%s AND u.vk_id>=9000000000 AND u.vk_id<10000000000
                    ORDER BY a.created_at DESC LIMIT 1''',(settings.recruitment_id,values['full_name']))
            actor=existing['vk_id'] if existing else 9_000_000_000+uuid4().int%1_000_000_000
            vk.synthetic.add(actor)
            app=await service.bot_repo.start_application(settings.recruitment_id,actor)
            # Only resume our own explicitly synthetic fixtures. Previously
            # stored answers remain immutable; existing real forms are untouched.
            await drain()
            await incoming(actor,'/start')
            for _ in range(24):
                dialog=await service.bot_repo.dialog_context(app['id'])
                if dialog['application']['status']=='submitted':
                    break
                current=next(q for q in dialog['questions'] if q['id']==dialog['application']['current_question_id'])
                await incoming(actor,values.get(current['target_slot'],CLARIFICATION))
            dialog=await service.bot_repo.dialog_context(app['id'])
            assert dialog['application']['status']=='submitted'
            assert len(dialog['answers'])>=19
            followups=[q for q in dialog['questions'] if q['source']!='base' and q.get('purpose')!='expectations']
            assert len(followups)<=3 and len({q['target_slot'] for q in followups})==len(followups)
            async with pools[1].connection() as c:
                evaluation=await one(c,'SELECT score,criteria_results,summary FROM gim.evaluations WHERE application_id=%s ORDER BY created_at DESC LIMIT 1',(app['id'],))
            if evaluation is None:
                raise RuntimeError('Live model did not produce a verified evaluation')
            report['applications'].append({'id':str(app['id']),'answers':len(dialog['answers']),
                'followups':[q['text'] for q in followups],'score':evaluation['score'],
                'criteria':len(evaluation['criteria_results']), 'expectations_checks':sum(q.get('purpose')=='expectations' for q in dialog['questions'])})
            print(json.dumps({'phase':'application_verified','application':str(app['id']),
                              'score':evaluation['score'],'answers':len(dialog['answers'])}),flush=True)
            await incoming(actor,'/status')
            await incoming(staff_vk_id,'/application '+str(app['id']))

        for command in ('/admin','/top','/stats','/export','/excel '+settings.excel_access_code):
            await incoming(staff_vk_id,command)
        bundle=await service.staff_repo.export_bundle(staff_vk_id,settings.recruitment_id)
        async with pools[1].connection() as c:
            files=await (await c.execute("SELECT id,storage_key,sha256,size_bytes,filters,status FROM gim.files WHERE recruitment_id=%s AND kind='xlsx' ORDER BY created_at DESC LIMIT 2",(settings.recruitment_id,))).fetchall()
        assert len(files)==2 and all(f['status']=='ready' for f in files)
        history=await vk.call('messages.getHistory',peer_id=staff_vk_id,count=30)
        docs=[a['doc'] for m in history['items'] for a in m.get('attachments',[]) if a['type']=='doc']
        async with httpx.AsyncClient(timeout=30,follow_redirects=True) as client:
            for meta in files:
                doc=next(d for d in docs if str(meta['id']) in d['title'])
                response=await client.get(doc['url'])
                response.raise_for_status()
                assert hashlib.sha256(response.content).hexdigest()==meta['sha256']
                wb=load_workbook(io.BytesIO(response.content),read_only=True)
                assert wb.sheetnames==['Анкеты','Ответы','Собеседование','Сводка','Оценка']
                assert wb['Анкеты'].max_row==len(bundle['Анкеты'])+1
                wb.close()
        report['xlsx_downloaded_and_verified']=len(files)
        report['code_export']=any('code_access' in f['filters'] for f in files)
        report['applications_total']=bundle['Сводка']['total']
        path=Path(output);path.parent.mkdir(parents=True,exist_ok=True)
        fd=os.open(path,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
        with os.fdopen(fd,'w') as f:json.dump(report,f,ensure_ascii=False,indent=2)
        print(json.dumps(report,ensure_ascii=False))


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--staff-vk-id',type=int,required=True)
    p.add_argument('--output',default='runtime/api-smoke.json')
    args=p.parse_args();load_dotenv()
    asyncio.run(smoke(args.staff_vk_id,args.output))


if __name__=='__main__':
    main()
