import copy
import io
import asyncio
import os
import unittest
from datetime import date
from types import SimpleNamespace
from uuid import uuid4
from unittest.mock import AsyncMock

from gim.shared_service import SharedService
from gim.shared_domain import parse_answer, select_followups
from gim.shared_xlsx import snapshot_xlsx

if os.name=='nt':
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())


class FakeRepository:
    def __init__(self, templates):
        self.pool = None
        self.app = dict(id=uuid4(), answers_revision=0, current_question_id=None,
                        status='draft', consent_at=None, consent_version=None)
        self.templates = templates
        self.questions, self.answers, self.events = [], [], {}
        self.contacts = []
        self.profile = None

    async def start_application(self, recruitment_id, verified_vk_id):
        return self.app.copy()

    async def dialog_context(self, app_id):
        return copy.deepcopy(dict(application=self.app, questions=self.questions, answers=self.answers,
                                  templates=self.templates, known_facts={}, direction_info={}, followup_policy={'max_followups':3}))

    async def prepare_question(self, app_id, **kw):
        key = kw['request_key']
        old = next((q for q in self.questions if q['request_key'] == key), None)
        if old:
            return old
        assert not self.app['current_question_id']
        t = next(t for t in self.templates if t['id'] == kw['template_id'])
        q = dict(id=uuid4(), text=t['text'], target_slot=t['target_slot'], template_id=t['id'],
                 source='base' if t['kind']=='base' else 'template', status='active', request_key=key,
                 purpose=kw.get('purpose','collect'))
        self.questions.append(q)
        self.app['current_question_id'] = q['id']
        return q

    async def record_answer(self, app_id, question_id, **kw):
        if kw['event_id'] in self.events:
            old = self.events[kw['event_id']]
            assert old['question_id'] == question_id and old['answer_text'] == kw['text']
            return old
        assert self.app['answers_revision'] == kw['expected_revision']
        self.app['answers_revision'] += 1
        a = dict(id=uuid4(), question_id=question_id, answer_text=kw['text'], value=kw['value'],
                 revision=self.app['answers_revision'])
        self.answers.append(a)
        self.events[kw['event_id']] = a
        self.app['current_question_id'] = None
        next(q for q in self.questions if q['id'] == question_id)['status'] = 'answered'
        return a

    async def set_consent(self, app_id, version):
        self.app.update(consent_at='now',consent_version=version)

    async def set_contacts(self, vk_id, **kw):
        assert self.app['consent_at']
        self.contacts.append(kw)

    async def set_profile(self, app_id, **kw):
        self.profile = kw

    async def submit(self, app_id):
        self.app['status'] = 'submitted'
        return self.app


class MemoryService(SharedService):
    def __init__(self, repo, settings):
        super().__init__(repo, repo, SimpleNamespace(export_bundle=AsyncMock(), request_export=AsyncMock()),
                         SimpleNamespace(send=AsyncMock(),upload=AsyncMock()),settings,SimpleNamespace())
        self.bindings, self.emitted = {}, {}
        self.crash_after_advance = False

    async def _bind(self, event, binding):
        return self.bindings.setdefault(str(event['id']), binding)

    async def _emit(self, event, peer, text, *, app_id=None, staff=False, suffix='reply'):
        self.emitted.setdefault((str(event['id']),suffix), dict(text=text,peer=peer,staff=staff))

    async def _staff(self, vk_id):
        return None


def template(code, pos, kind='base', directions=None):
    return dict(id=uuid4(),code=code,target_slot=code,text='Вопрос '+code,position=pos,
                kind=kind,required=kind=='base',directions=directions or [],validation={})


class SharedServiceTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.repo = FakeRepository([template(c,i) for i,c in enumerate(('consent','full_name','email','birth_date','directions'))])
        self.settings = SimpleNamespace(group_id=17,recruitment_id=uuid4(),consent_version='v1',max_followups=3,max_answer_length=2000)
        self.s = MemoryService(self.repo,self.settings)

    def event(self, text, **extra):
        normalized = dict(type='message_new',group_id=17,event_id=str(uuid4()),object={'message':dict(from_id=42,peer_id=42,text=text,**extra)})
        return dict(id=uuid4(),lease_token=uuid4(),payload=dict(event=normalized,group_id=17,serial_key='17:42'))

    async def test_persisted_question_precedes_external_send(self):
        await self.s.incoming(self.event('/start'))
        self.assertEqual(self.repo.questions[0]['target_slot'],'consent')
        self.s.vk.send.assert_not_awaited()
        self.assertEqual(self.repo.answers,[])

    async def test_replay_after_advance_never_answers_next_question(self):
        await self.s.incoming(self.event('/start'))
        original = self.event('Согласен')
        await self.s.incoming(original)
        current = self.repo.app['current_question_id']
        await self.s.incoming(original)
        self.assertEqual(len(self.repo.answers),1)
        self.assertEqual(self.repo.app['current_question_id'],current)
        self.assertEqual(self.repo.questions[-1]['target_slot'],'full_name')

    async def test_raw_is_never_truncated_and_contacts_require_consent(self):
        await self.s.incoming(self.event('/start'))
        await self.s.incoming(self.event('Согласен'))
        raw = '=synthetic ' + 'x'*5000
        await self.s.incoming(self.event(raw))
        self.assertEqual(self.repo.answers[-1]['answer_text'],raw)
        self.assertEqual(self.repo.contacts,[])
        await self.s.incoming(self.event('synthetic@example.test'))
        self.assertEqual(self.repo.contacts[-1]['full_name'],raw)

    async def test_consent_denial_does_not_collect_contacts(self):
        await self.s.incoming(self.event('/start'))
        await self.s.incoming(self.event('Нет'))
        self.assertEqual(self.repo.answers[-1]['answer_text'],'Нет')
        self.assertEqual(self.repo.contacts,[])
        self.assertIsNone(self.repo.app['consent_at'])
        self.assertEqual(len(self.repo.questions),2)
        self.assertEqual(self.repo.questions[-1]['target_slot'],'consent')

    async def test_group_and_payload_actor_are_not_authorization(self):
        event = self.event('/export')
        event['payload']['actor_vk'] = 1
        event['payload']['role'] = 'admin'
        await self.s.incoming(event)
        self.s.staff_repo.request_export.assert_not_awaited()
        event['payload']['event']['group_id'] = 99
        before = len(self.s.emitted)
        await self.s.incoming(event)
        self.assertEqual(len(self.s.emitted),before)

    async def test_completed_dialog_sets_provenance_before_submit(self):
        for text in ['/start','да','Synthetic','synthetic@example.test','01.01.2000','не определился']:
            await self.s.incoming(self.event(text))
        self.assertEqual(self.repo.app['status'],'submitted')
        self.assertEqual(self.repo.profile['birth_date'],date(2000,1,1))
        self.assertTrue(self.repo.profile['undecided'])
        self.assertEqual(self.repo.profile['birth_date_answer_id'],self.repo.answers[-2]['id'])
        self.assertTrue(all(x['answer_ids'] for x in self.repo.profile['facts'].values()))

    async def test_missing_key_keeps_analysis_retryable(self):
        from gim.llm import ModelUnavailable
        with self.assertRaises(ModelUnavailable):
            await self.s.analyze(dict(snapshot={},application_id=uuid4()))

    async def test_followup_provider_outage_does_not_block_submission(self):
        self.s.settings.llm_api_key='synthetic'
        self.s.llm.complete=AsyncMock(side_effect=TimeoutError)
        for text in ['/start','да','Synthetic','synthetic@example.test','01.01.2000','не определился']:
            await self.s.incoming(self.event(text))
        self.assertEqual(self.repo.app['status'],'submitted')


class DomainTests(unittest.TestCase):
    def test_age_and_directions(self):
        self.assertIsNotNone(parse_answer('birth_date','01.01.2020')[1])
        self.assertEqual(parse_answer('directions','организационное, медиа')[0]['selected_directions'],['organizational','media'])
        self.assertIsNotNone(parse_answer('directions','что угодно')[1])
        self.assertIsNotNone(parse_answer('consent','не согласен')[1])

    def test_followups_are_from_library_and_at_most_three(self):
        ts = [template('experience',0)] + [template('followup_'+str(i),i,'followup',['media']) for i in range(5)]
        selected = select_followups(ts, [], ['media'], 9)
        self.assertLessEqual(len(selected),3)
        self.assertTrue(all(t in ts for t in selected))

    def test_followup_direction_does_not_invent_past_experience(self):
        interest=template('collections_interest',1,'followup',['collections'])
        interest['text']='Что вам интересно в коллекциях?'
        past=template('collections_past',2,'followup',['collections'])
        past['text']='С какими коллекциями вы работали раньше?'
        wrong=template('media_past',3,'followup',['media'])
        wrong['text']='Какие видео вы делали раньше?'
        generic=template('past_generic',4,'followup',[])
        generic['text']='В каких проектах вы участвовали раньше?'
        bank=[interest,past,wrong,generic]
        self.assertEqual(select_followups(bank,[],['collections'],3,
                         {'experience':{'answer_text':'нет'}}),[])
        self.assertEqual(select_followups(bank,[],['collections'],3,
                         {'experience':{'answer_text':'Работал с музейными коллекциями'}}),[past,generic])

    def test_xlsx_preserves_full_text_and_no_formula(self):
        from openpyxl import load_workbook
        raw = '=SUM(A1:A2)' + 'x'*70000
        bundle = {'Анкеты':[], 'Ответы':[dict(application_id='app',question_id='q',answer_id='a',answer_text=raw)],
                  'Собеседование':[], 'Сводка':dict(total=1,scored=0,average_score=None)}
        wb = load_workbook(io.BytesIO(snapshot_xlsx(bundle)))
        self.assertEqual(wb.sheetnames,['Анкеты','Ответы','Собеседование','Сводка'])
        rows = list(wb['Ответы'].values)
        ix = rows[0].index('answer_text')
        self.assertEqual(''.join(r[ix] for r in rows[1:]),raw)
        self.assertTrue(all(c.data_type!='f' for row in wb['Ответы'] for c in row))

class ConnectionScope:
    def __init__(self, value): self.value=value
    async def __aenter__(self): return self.value
    async def __aexit__(self,*args): pass


class FakeSQL:
    def __init__(self, rows): self.rows=list(rows); self.calls=[]
    def connection(self): return ConnectionScope(self)
    def transaction(self): return ConnectionScope(self)
    async def execute(self, sql, args=()):
        self.calls.append((sql,args))
        return self
    async def fetchone(self): return self.rows.pop(0)


class BoundaryTests(unittest.IsolatedAsyncioTestCase):
    def service(self,pool):
        repo=SimpleNamespace(pool=pool)
        return SharedService(repo,repo,SimpleNamespace(staff_file=AsyncMock()),
              SimpleNamespace(send=AsyncMock(),upload=AsyncMock(return_value='doc_synthetic')),
              SimpleNamespace(group_id=17,recruitment_id=uuid4()),
              SimpleNamespace(upload=AsyncMock(),read=AsyncMock()))

    async def test_binding_uses_current_lease_and_keeps_original(self):
        original=dict(app_id='a',question_id='q1',revision=0)
        sql=FakeSQL([{'binding':original}])
        service=self.service(sql)
        event=dict(id=uuid4(),lease_token=uuid4())
        bound=await service._bind(event,dict(app_id='a',question_id='q2',revision=1))
        self.assertEqual(bound,original)
        statement,args=sql.calls[0]
        self.assertIn('COALESCE',statement)
        self.assertIn('lease_token=%s',statement)
        self.assertEqual(args[-2:],(event['id'],event['lease_token']))

    async def test_expired_lease_cannot_record_binding(self):
        from gim_db.repository import Conflict
        service=self.service(FakeSQL([None]))
        with self.assertRaises(Conflict):
            await service._bind(dict(id=uuid4(),lease_token=uuid4()),{})

    async def test_reply_is_an_outbox_intention(self):
        sql=FakeSQL([None])
        service=self.service(sql)
        event=dict(id=uuid4())
        await service._emit(event,42,'synthetic')
        query,args=sql.calls[0]
        self.assertIn('ON CONFLICT(dedupe_key) DO NOTHING',query)
        self.assertEqual(args[1],'vk.send')
        service.vk.send.assert_not_awaited()

    async def test_staff_send_rechecks_membership(self):
        service=self.service(FakeSQL([None]))
        await service.send(dict(payload=dict(group_id=17,peer_id=42,text='synthetic',staff=True),vk_random_id=99))
        service.vk.send.assert_not_awaited()

    async def test_export_ready_reuses_immutable_file_and_checks_acl_twice(self):
        import hashlib
        data=b'synthetic_xlsx'
        meta=dict(id=uuid4(),kind='xlsx',status='ready',requested_by_staff_id=uuid4(),
                  size_bytes=len(data),sha256=hashlib.sha256(data).hexdigest())
        service=self.service(FakeSQL([meta,dict(vk_id=42,group_id=17)]))
        service.storage.read.return_value=data
        service.files.finish=AsyncMock()
        await service.export(dict(file_id=meta['id'],payload={},vk_random_id=99))
        service.storage.upload.assert_not_awaited()
        service.files.finish.assert_not_awaited()
        self.assertEqual(service.staff_repo.staff_file.await_count,2)
        self.assertEqual(service.vk.send.await_args.args[2],99)

    async def test_export_uses_snapshot_then_finishes_before_vk(self):
        meta=dict(id=uuid4(),kind='xlsx',status='pending',requested_by_staff_id=uuid4())
        service=self.service(FakeSQL([meta,dict(vk_id=42,group_id=17)]))
        order=[]
        async def upload(meta,data): order.append('storage')
        async def finish(fid,**kw): order.append('finish'); return {**meta,'status':'ready',**kw}
        async def vk_upload(*args): order.append('vk'); return 'doc_synthetic'
        service.storage.upload.side_effect=upload
        service.files.finish=AsyncMock(side_effect=finish)
        service.vk.upload.side_effect=vk_upload
        await service.export(dict(file_id=meta['id'],payload=dict(snapshot={'Анкеты':[],'Ответы':[],'Собеседование':[],'Сводка':{}}),vk_random_id=101))
        self.assertEqual(order,['storage','finish','vk'])
        service.storage.read.assert_not_awaited()

    async def test_export_rejects_incomplete_snapshot_before_writing_file(self):
        meta=dict(id=uuid4(),kind='xlsx',status='pending',requested_by_staff_id=uuid4())
        service=self.service(FakeSQL([meta,dict(vk_id=42,group_id=17)]))
        service.files.finish=AsyncMock()
        with self.assertRaises(ValueError):
            await service.export(dict(file_id=meta['id'],payload=dict(snapshot={'unexpected':[]}),vk_random_id=101))
        service.storage.upload.assert_not_awaited()
        service.files.finish.assert_not_awaited()


class MoreDialogTests(SharedServiceTests):
    async def test_consent_denial_then_agreement_is_not_dropped(self):
        await self.s.incoming(self.event('/start'))
        first=self.repo.app['current_question_id']
        await self.s.incoming(self.event('Нет'))
        repeated=self.repo.app['current_question_id']
        self.assertIsNotNone(repeated)
        self.assertNotEqual(repeated,first)
        self.assertEqual(self.repo.questions[-1]['target_slot'],'consent')
        self.assertIsNone(self.repo.app['consent_at'])
        await self.s.incoming(self.event('Согласен'))
        self.assertEqual([a['answer_text'] for a in self.repo.answers],['Нет','Согласен'])
        self.assertIsNotNone(self.repo.app['consent_at'])
        self.assertEqual(self.repo.questions[-1]['target_slot'],'full_name')

    async def test_invalid_direction_then_valid_choice_is_not_dropped(self):
        for text in ['/start','да','Synthetic','synthetic@example.test','01.01.2000','неизвестный выбор']:
            await self.s.incoming(self.event(text))
        self.assertEqual(self.repo.app['status'],'draft')
        self.assertIsNotNone(self.repo.app['current_question_id'])
        self.assertEqual(self.repo.questions[-1]['target_slot'],'directions')
        self.assertEqual(self.repo.answers[-1]['answer_text'],'неизвестный выбор')
        await self.s.incoming(self.event('медиа'))
        self.assertEqual(self.repo.app['status'],'submitted')
        self.assertEqual(self.repo.profile['selected_directions'],['media'])
        self.assertEqual(self.repo.answers[-1]['answer_text'],'медиа')

    async def test_invalid_date_is_preserved_and_corrected_without_losing_history(self):
        for text in ['/start','да','Synthetic','synthetic@example.test','01.01.2020']:
            await self.s.incoming(self.event(text))
        self.assertEqual(self.repo.answers[-1]['answer_text'],'01.01.2020')
        self.assertEqual(self.repo.app['status'],'draft')
        self.assertEqual(self.repo.questions[-1]['target_slot'],'birth_date')
        await self.s.incoming(self.event('01.01.2000'))
        self.assertEqual(self.repo.questions[-1]['target_slot'],'directions')
        self.assertEqual(sum(a['answer_text'].startswith('01.01.') for a in self.repo.answers),2)

    async def test_crash_after_answer_before_next_question_recovers_same_binding(self):
        await self.s.incoming(self.event('/start'))
        event=self.event('да')
        advance=self.s._advance
        self.s._advance=AsyncMock(side_effect=RuntimeError('synthetic_crash'))
        with self.assertRaises(RuntimeError):
            await self.s.incoming(event)
        self.s._advance=advance
        event['payload']['binding']=self.s.bindings[str(event['id'])]
        await self.s.incoming(event)
        self.assertEqual(len(self.repo.answers),1)
        self.assertEqual(self.repo.questions[-1]['target_slot'],'full_name')


class SnapshotRepeatTests(unittest.TestCase):
    def test_export_bytes_are_identical_after_time_changes(self):
        import time
        snapshot={'Анкеты':[],'Ответы':[],'Собеседование':[],'Сводка':{'total':0}}
        before=snapshot_xlsx(snapshot)
        time.sleep(1.1)
        self.assertEqual(before,snapshot_xlsx(snapshot))

# Explicit local-only integration: never infer credentials or read .env.
@unittest.skipUnless(os.environ.get('TEST_DATABASE_URL'),'TEST_DATABASE_URL not provided')
class SharedDatabaseIntegration(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        import psycopg
        from psycopg.conninfo import conninfo_to_dict
        from psycopg.rows import dict_row
        from psycopg.types.json import Jsonb
        from psycopg_pool import AsyncConnectionPool
        from gim_db.repository import Repository
        from gim_db.outbox import SQLQueue
        dsn=os.environ['TEST_DATABASE_URL']
        if conninfo_to_dict(dsn).get('host') not in {'localhost','127.0.0.1','::1'}:
            self.skipTest('Integration fixture permitted only on an explicit local test server')
        self.group=uuid4().int%1000000000+100
        self.candidate=self.group+2000000000
        self.admin=self.group+3000000000
        self.codes=['consent','full_name','email','vk_url','birth_date','current_status','university',
                    'university_other','faculty','course','interests','experience','languages','motivation',
                    'expectations','available_time','hobbies','directions','discovery']
        self.owner=await psycopg.AsyncConnection.connect(dsn,autocommit=True,row_factory=dict_row)
        async with self.owner.transaction():
            row=await (await self.owner.execute('''INSERT INTO gim.recruitments(group_id,title,questionnaire_version,
                rubric_version,rubric,followup_policy) VALUES(%s,'Synthetic shared integration','test','test',%s,%s) RETURNING id''',
                (self.group,Jsonb({'criteria':[{'code':'synthetic','description':'Synthetic only','max_points':3}],
                    'approved_by':'synthetic-test-only','scoring_method':'sum_round_half_up',
                    'scoring_rule':'Сумма, округление до целого (0.5 вверх)'}),Jsonb({'max_followups':3})))).fetchone()
            self.rid=row['id']
            for i,code in enumerate(self.codes,1):
                await self.owner.execute('''INSERT INTO gim.question_templates(recruitment_id,code,kind,text,position,required,target_slot)
                    VALUES(%s,%s,'base',%s,%s,true,%s)''',(self.rid,code,'Synthetic '+code,i,code))
            for i in range(1,5):
                await self.owner.execute('''INSERT INTO gim.question_templates(recruitment_id,code,kind,text,position,target_slot,directions)
                    VALUES(%s,%s,'followup',%s,%s,%s,ARRAY['media'])''',
                    (self.rid,'followup_'+str(i),'Какие работы вы делали раньше? '+str(i),30+i,'past_media_'+str(i)))
            await self.owner.execute("UPDATE gim.recruitments SET status='active',opened_at=now() WHERE id=%s",(self.rid,))
            user=await (await self.owner.execute('INSERT INTO gim.users(vk_id) VALUES(%s) RETURNING id',(self.admin,))).fetchone()
            await self.owner.execute("INSERT INTO gim.staff_members(user_id,group_id,role) VALUES(%s,%s,'admin')",(user['id'],self.group))
        self.pools=[]
        for role in ('gim_bot','gim_worker','gim_staff'):
            pool=AsyncConnectionPool(dsn,min_size=1,max_size=3,open=False,kwargs={'autocommit':True,'row_factory':dict_row,
                       'prepare_threshold':None,'options':'-c role='+role})
            await pool.open(wait=True)
            self.pools.append(pool)
        self.bot,self.worker,self.staff=[Repository(p) for p in self.pools]
        self.objects={}
        async def upload(meta,data):
            old=self.objects.setdefault(meta['storage_key'],data)
            assert old==data
        async def read(meta): return self.objects[meta['storage_key']]
        self.vk=SimpleNamespace(send=AsyncMock(),upload=AsyncMock(return_value='doc_synthetic'))
        self.service=SharedService(self.bot,self.worker,self.staff,self.vk,
                 SimpleNamespace(group_id=self.group,recruitment_id=self.rid,consent_version='test-only',max_followups=3,max_answer_length=2000),
                 SimpleNamespace(upload=upload,read=read))
        self.queue=SQLQueue(self.worker.pool,self.group,lease_seconds=30)
        self.ctx=dict(repository=self.worker,handlers=self.service.handlers)

    async def asyncTearDown(self):
        if hasattr(self,'pools'):
            for pool in self.pools: await pool.close()
        if hasattr(self,'owner'): await self.owner.close()

    async def enqueue(self,text,actor=None):
        from gim_db.repository import emit
        actor=actor or self.candidate
        event_id=str(uuid4())
        normalized=dict(type='message_new',group_id=self.group,event_id=event_id,
                        object={'message':dict(from_id=actor,peer_id=actor,text=text)})
        async with self.worker.pool.connection() as c:
            return await emit(c,None,'vk.incoming','synthetic:'+event_id,
                {'event':normalized,'group_id':self.group,'serial_key':f'{self.group}:{actor}'})

    async def drain(self):
        from gim_db.outbox import run_once
        for _ in range(20):
            if not await run_once(self.ctx,self.queue,limit=20): break
        pending=await (await self.owner.execute('''SELECT e.topic,e.last_error_code FROM gim.outbox_events e
           LEFT JOIN gim.applications a ON a.id=e.application_id LEFT JOIN gim.recruitments r ON r.id=a.recruitment_id
           LEFT JOIN gim.files f ON f.id=e.file_id LEFT JOIN gim.recruitments fr ON fr.id=f.recruitment_id
           WHERE COALESCE(r.group_id::text,fr.group_id::text,e.payload->>'group_id')=%s
           AND e.completed_at IS NULL''',(str(self.group),))).fetchall()
        if not getattr(self.service.settings,'llm_api_key',''):
            self.assertTrue(all(e['topic']=='application.analyze' and e['last_error_code']=='ModelUnavailable' for e in pending))
        else:
            self.assertEqual(pending,[])

    async def test_real_private_roles_replay_followups_submit_export(self):
        from gim_db.outbox import run_once
        await self.enqueue('/start')
        await self.drain()
        await self.enqueue('Нет')
        await self.drain()
        consent=await self.enqueue('Согласен')
        claimed=await self.queue.claim(['vk.incoming'])
        self.assertEqual(claimed['id'],consent['id'])
        # Crash after record+advance but before durable worker acknowledgement.
        await self.service.incoming(claimed)
        await self.queue.finish(claimed['id'],claimed['lease_token'],False,'synthetic_crash')
        await self.owner.execute('UPDATE gim.outbox_events SET available_at=now() WHERE id=%s',(claimed['id'],))
        await self.drain()
        app=await self.bot.start_application(self.rid,self.candidate)
        ctx=await self.bot.dialog_context(app['id'])
        self.assertEqual(len(ctx['answers']),2)
        self.assertEqual([a['answer_text'] for a in ctx['answers']],['Нет','Согласен'])
        self.assertEqual(next(q for q in ctx['questions'] if q['id']==ctx['application']['current_question_id'])['target_slot'],'full_name')
        raw='=synthetic ' + 'x'*70000
        values={'full_name':'Synthetic Person','email':'synthetic@example.test','vk_url':'https://vk.com/synthetic',
                'birth_date':'01.01.2000','directions':'медиа','experience':raw}
        for code in self.codes[1:]:
            await self.enqueue(values.get(code,'не относится'))
            await self.drain()
        for _ in range(3):
            await self.enqueue('Synthetic past work')
            await self.drain()
        ctx=await self.bot.dialog_context(app['id'])
        self.assertEqual(ctx['application']['status'],'submitted')
        self.assertEqual(len(ctx['answers']),23)
        self.assertEqual(sum(q['source']!='base' for q in ctx['questions']),3)
        self.assertEqual(ctx['application']['selected_directions'],['media'])
        count=await (await self.owner.execute('SELECT count(*) AS n FROM gim.evaluations WHERE application_id=%s',(app['id'],))).fetchone()
        self.assertEqual(count['n'],0)
        disabled=await (await self.owner.execute("SELECT last_error_code FROM gim.outbox_events WHERE application_id=%s AND topic='application.analyze'",(app['id'],))).fetchone()
        self.assertEqual(disabled['last_error_code'],'ModelUnavailable')
        self.service.settings.llm_api_key='synthetic-never-sent'
        self.service.settings.llm_model='synthetic-model'
        answer=next(a for a in ctx['answers'] if a['answer_text']==raw)
        self.service.llm.complete=AsyncMock(return_value={
            'score':2,'criteria':[{'code':'synthetic','points':2,'reason':'Synthetic only',
            'evidence':[{'source_type':'answer','source_id':str(answer['id']),'quote':'=synthetic'}]}],
            'summary':'Synthetic summary','directions':['media'],'interview_questions':['Какова ваша роль?']})
        await self.owner.execute("UPDATE gim.outbox_events SET available_at=now() WHERE application_id=%s AND topic='application.analyze'",(app['id'],))
        await self.drain()
        evaluation=await (await self.owner.execute('SELECT * FROM gim.evaluations WHERE application_id=%s',(app['id'],))).fetchone()
        self.assertEqual(evaluation['score'],2)
        self.assertEqual(evaluation['criteria_results'][0]['evidence'][0]['quote'],'=synthetic')
        await self.enqueue('/export',self.admin)
        await self.drain()
        meta=await (await self.owner.execute("SELECT * FROM gim.files WHERE recruitment_id=%s AND kind='xlsx'",(self.rid,))).fetchone()
        self.assertEqual(meta['status'],'ready')
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(self.objects[meta['storage_key']]))
        rows=list(wb['Ответы'].values)
        cols=rows[0]
        slot=cols.index('target_slot'); answer=cols.index('answer_text')
        # Each continuation carries its answer_id; field slot itself is first-row only.
        aid=next(r[cols.index('answer_id')] for r in rows[1:] if r[slot]=='experience')
        reconstructed=''.join(r[answer] or '' for r in rows[1:] if r[cols.index('answer_id')]==aid)
        self.assertEqual(reconstructed,raw)
        exported=list(wb['Анкеты'].values)
        self.assertEqual(exported[1][exported[0].index('score')],'2')
        self.vk.upload.assert_awaited_once()

class WorkbookEncodingTests(unittest.TestCase):
    def test_xml_control_character_is_recoverable(self):
        import json
        from openpyxl import load_workbook
        raw='synthetic\x00control and literal \\u0000'
        wb=load_workbook(io.BytesIO(snapshot_xlsx({'Ответы':[dict(answer_id='a',answer_text=raw)]})))
        rows=list(wb['Ответы'].values)
        row=dict(zip(rows[0],rows[1]))
        self.assertEqual(json.loads(row['answer_text']),raw)
        self.assertEqual(json.loads(row['field_encodings']),{'answer_text':'json_string'})
