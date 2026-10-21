import unittest,io
from openpyxl import load_workbook
from gim.dashboard import dashboard
class DashboardTests(unittest.TestCase):
 def row(self):
  return {'application_id':'a','score':2,'criteria_results':[{'code':c,'points':p,'needs_review':False,'calculation':{'kind':'binary_50_35_15','raw':1.98} if i==0 else {}} for i,(c,p) in enumerate([('motivation',1.125),('experience',.63),('availability',.225)])]}
 def test_three_stored_dimensions_without_model_call(self):
  text=dashboard(self.row())
  self.assertIn('Мотивация',text);self.assertIn('Опыт',text);self.assertIn('Занятость',text)
  self.assertIn('1.125 / 1.5',text)
  self.assertIn('не вероятность',text)
 def test_unknown_is_not_shown_as_zero(self):
  r=self.row();r['criteria_results'][2]['needs_review']=True
  self.assertIn('недостаточно данных',dashboard(r))
 def test_no_evaluation_never_invents_graph(self):
  self.assertIn('не готов',dashboard({'score':None}))
 def test_dashboard_is_in_excel(self):
  from gim.shared_xlsx import snapshot_xlsx
  w=load_workbook(io.BytesIO(snapshot_xlsx({'Анкеты':[self.row()]})))
  h=list(next(w['Анкеты'].values));self.assertIn('Профиль кандидата',h)
  proof=w['Оценка'];self.assertEqual(len(proof.conditional_formatting),1)
  columns=list(next(proof.values));self.assertEqual(proof.cell(2,columns.index('Основания / максимум группы')+1).value,.75)
