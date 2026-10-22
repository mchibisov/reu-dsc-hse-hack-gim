import io,unittest
from openpyxl import load_workbook
from gim.shared_service import question_text
from gim.shared_xlsx import snapshot_xlsx

class FinalFeatures(unittest.TestCase):
    def test_base_question_shows_progress(self):
        text=question_text({'text':'Ваш опыт?','target_slot':'experience','position':12,'source':'base'},42)
        self.assertIn('12 из 18',text)

    def test_faculty_prompt_has_requested_name(self):
        text=question_text({'text':'Старое название','target_slot':'faculty','position':8,'source':'base'},42)
        self.assertIn('Факультет, направление',text)
        self.assertNotIn('Старое название',text)

    def test_export_is_a_readable_working_list_with_numeric_scores(self):
        bundle={'Анкеты':[{'application_id':'app','full_name':'Тест','summary':'Содержательная сводка','score':2,'selected_directions':['media']}],
                'Ответы':[],'Собеседование':[],'Сводка':{'total':1}}
        w=load_workbook(io.BytesIO(snapshot_xlsx(bundle)));s=w['Анкеты'];headers=list(next(s.values))
        self.assertEqual(headers[0],'ФИО')
        self.assertEqual(s.cell(2,headers.index('Балл 0–3')+1).value,2)
        self.assertEqual(s.cell(2,headers.index('Выбранные направления')+1).value,'Медиа')
        self.assertTrue(s.cell(2,headers.index('AI-сводка')+1).alignment.wrap_text)
        self.assertGreater(s.column_dimensions['A'].width,20)

    def test_expectations_check_prevents_repeating_alternative_question(self):
        from gim.shared_domain import select_followups
        template={'position':1,'id':'t','target_slot':'alternatives','text':'Готовы ли рассмотреть другие направления?','kind':'followup','directions':[]}
        asked=[{'source':'generated','purpose':'expectations','target_slot':'expectations_match','text':'Подходят ли задачи и интересны ли альтернативы?'}]
        self.assertEqual(select_followups([template],asked,['collections'],3,{}),[])
