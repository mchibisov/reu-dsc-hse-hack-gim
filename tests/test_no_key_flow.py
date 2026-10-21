"""Один сквозной сценарий без ключа LLM; VK API имитируется на уровне HTTP."""
import json
import os
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs

import httpx
from openpyxl import load_workbook
from test_bot import event
from test_dialog import ANSWERS

from gim_bot.analysis import Provider
from gim_bot.bot import Bot
from gim_bot.config import Settings
from gim_bot.vk import VKClient
from gim_bot.worker import run_once


async def test_complete_vk_questionnaire_and_staff_excel_without_llm_key(store):
    settings = Settings(vk_token='synthetic-vk-token', vk_group_id=1, admin_ids='99',
                        llm_api_key='', llm_model='', gateway_url='', internal_api_key='')
    sent, uploads, methods, llm_calls = [], [], [], []

    def forbidden_llm(request):
        llm_calls.append(request)
        raise AssertionError('Без ключа модели внешние LLM-запросы запрещены')

    def vk_api(request):
        method = request.url.path.rsplit('/', 1)[-1]
        methods.append(method)
        if request.url.host == 'upload.invalid':
            mime = BytesParser(policy=policy.default).parsebytes(
                b'Content-Type: ' + request.headers['content-type'].encode() + b'\r\n\r\n' + request.content)
            uploads.extend(part.get_payload(decode=True) for part in mime.iter_parts() if part.get_filename())
            body = {'file': 'synthetic-upload'}
        elif method == 'docs.getMessagesUploadServer':
            body = {'response': {'upload_url': 'https://upload.invalid/xlsx'}}
        elif method == 'docs.save':
            body = {'response': {'type': 'doc', 'doc': {'owner_id': 1, 'id': 2}}}
        elif method == 'messages.send':
            sent.append(parse_qs(request.content.decode()))
            body = {'response': len(sent)}
        else:
            raise AssertionError(f'Неожиданный VK метод: {method}')
        return httpx.Response(200, json=body)

    vk = VKClient(settings, transport=httpx.MockTransport(vk_api), rate_interval=0)
    provider = Provider(settings, transport=httpx.MockTransport(forbidden_llm))
    bot = Bot(settings, store, vk, provider)
    context = {'store': store, 'settings': settings, 'vk': vk, 'gateway': provider}
    try:
        await bot.process_update(event('/start', 1))
        await bot.process_update(event('Согласен', 2))
        for index in range(40):
            app = await store.get_by_user(123)
            if app.status == 'completed':
                break
            key = await store.current_key(123)
            await bot.process_update(event(ANSWERS.get(key, 'Снимал музейные события в команде.'), index + 3))
        await run_once(context)
        app = await store.get_by_user(123)
        assert app.status == 'completed'
        assert app.analysis_status == 'no_analytics' and app.score is None and app.analysis is None
        assert app.answers['experience'] == ANSWERS['experience']
        assert app.answers['directions'] == ANSWERS['directions']
        assert app.followups and all(item['answer'] and item['source'] == 'museum' for item in app.followups)
        candidate_messages = [item['message'][0] for item in sent if item['peer_id'] == ['123']]
        assert 'получена' in candidate_messages[-1].lower()
        assert 'балл' not in candidate_messages[-1].lower()
        assert any('+1 заполненная заявка' in item['message'][0] for item in sent if item['peer_id'] == ['99'])
        for number, command in enumerate(('/admin', '/stats', f'/candidate {app.id}', '/export'), 100):
            await bot.process_update(event(command, number, uid=99))
        await run_once(context)
        assert len(uploads) == 1
        workbook = load_workbook(BytesIO(uploads[0]))
        assert workbook.sheetnames == ['Анкеты', 'Сводка']
        rows = list(workbook['Анкеты'].values)
        assert len(rows) == 2
        assert ANSWERS['experience'] in rows[1] and ANSWERS['email'] in rows[1]
        assert any(item.get('attachment') == ['doc1_2'] for item in sent)
        assert not await store.pending_exports()
        await bot.process_update(event('/export', 103, uid=99))
        await run_once(context)
        assert len(uploads) == 1
        assert not llm_calls
        report = os.environ.get('GIM_NO_KEY_REPORT')
        if report:
            Path(report).write_text(json.dumps({
                'scope': 'synthetic data; real application classes; mocked VK HTTP',
                'completed': True, 'raw_fields_saved': len(app.answers),
                'answered_followups': len(app.followups), 'analysis_status': app.analysis_status,
                'score': app.score, 'llm_http_calls': len(llm_calls), 'xlsx_valid': True,
                'xlsx_candidate_rows': len(rows) - 1, 'uploads_after_replay': len(uploads),
                'vk_methods': sorted(set(methods)),
            }, ensure_ascii=False, indent=2), encoding='utf-8')
    finally:
        await vk.close()
        await provider.close()
