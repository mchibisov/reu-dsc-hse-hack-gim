from io import BytesIO

from openpyxl import load_workbook
from test_dialog import finish

from gim_bot.admin import Admin
from gim_bot.config import Settings
from gim_bot.export import build_excel


async def test_admin_denies_unknown_id_and_group_chat(store):
    admin = Admin(store, Settings(admin_ids='999'))
    assert await admin.handle(123, 123, '/top', '1') is None
    assert await admin.handle(999, 2000000001, '/top', '2') is None
    reply = await admin.handle(999, 999, '/stats', '3')
    assert 'заверш' in reply.lower()


async def test_admin_export_request_is_durable_and_idempotent(store):
    await finish(store)
    admin = Admin(store, Settings(admin_ids='999'))
    await admin.handle(999, 999, '/export', 'export-event')
    await admin.handle(999, 999, '/export', 'export-event')
    requests = await store.pending_exports()
    assert len(requests) == 1
    assert requests[0].admin_id == 999


async def test_excel_contains_answers_and_neutralizes_formula(store):
    app = await finish(store)
    app.answers = {**app.answers, 'hobbies': '=HYPERLINK("https://example.org")'}
    data = build_excel([app], 1)
    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == ['Анкеты', 'Сводка']
    worksheet = workbook['Анкеты']
    assert worksheet.auto_filter.ref
    assert worksheet.freeze_panes == 'A2'
    assert not any(cell.data_type == 'f' for row in worksheet for cell in row)
    assert any('HYPERLINK' in str(cell.value) for row in worksheet for cell in row)


async def test_admin_has_no_edit_commands(store):
    admin = Admin(store, Settings(admin_ids='999'))
    reply = await admin.handle(999, 999, '/score 1 3', 'event')
    assert 'не поддерживается' in reply
