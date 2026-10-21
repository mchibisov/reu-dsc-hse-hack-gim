import os
import tempfile
import unittest
import uuid
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock
from unittest.mock import patch

from gim.db import Database
from gim.service import Service


@unittest.skipUnless(os.environ.get('TEST_DATABASE_URL'), 'TEST_DATABASE_URL not set')
class ServiceTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.db = Database(os.environ['TEST_DATABASE_URL'])
        await self.db.open()
        async with self.db.pool.connection() as c:
            await c.execute('TRUNCATE users,inbox_events,jobs CASCADE')
        self.group = uuid.uuid4().int % 100000000
        self.db.group_id = self.group
        self.vk_id = uuid.uuid4().int % 100000000 + 100000000
        self.settings = SimpleNamespace(group_id=self.group, campaign='test', consent_text='Тестовое согласие',
                                        storage_dir=tempfile.mkdtemp(prefix='gim-test-'), scoring_url='', scoring_token='',
                                        followup_url='https://questions.test/next', service_token='', max_followups=3)
        from pathlib import Path
        rubric_path = Path(self.settings.storage_dir) / 'rubric.json'
        rubric_path.write_text(json.dumps({'version': 'team-v1', 'criteria': [
            {'code': 'experience', 'description': 'Критерий команды для теста', 'max_score': 3}]}))
        self.settings.rubric_path = str(rubric_path)
        self.auto_ai = True
        self.vk = SimpleNamespace(send=AsyncMock(), upload=AsyncMock(return_value='doc1_2'))
        self.s = Service(self.db, self.vk, self.settings)

    async def asyncTearDown(self):
        await self.db.close()

    async def process(self, text, attachments=None):
        event = {'type': 'message_new', 'group_id': self.group, 'event_id': str(uuid.uuid4()),
                 'object': {'message': {'from_id': self.vk_id, 'peer_id': self.vk_id, 'text': text, 'attachments': attachments or []}}}
        await self.db.ingest(event)
        # Claim only this test's earliest dialog through ordinary queue semantics.
        job = await self.db.claim(kinds=['dialog'])
        self.assertIsNotNone(job)
        await self.s.dialog(job)
        if self.auto_ai:
            ai_job = await self.db.claim(kinds=['prepare_followups'])
            if ai_job:
                import httpx
                import json
                async def endpoint(request):
                    payload = json.loads(request.content)
                    if any(a['question_code'].startswith('followup_') for a in payload['answers']):
                        result = {'action': 'finish'}
                    else:
                        result = {'action': 'ask', 'direction': payload['selected_directions'][0],
                                  'question': 'Расскажите, как вы проверяли точность записей в конкретном проекте?'}
                    return httpx.Response(200, json=result)
                factory = httpx.AsyncClient
                with patch('gim.service.httpx.AsyncClient', lambda **kwargs: factory(transport=httpx.MockTransport(endpoint), **kwargs)):
                    await self.s.prepare_followups(ai_job)
                await self.db.finish(ai_job)
        return job

    async def test_answers_and_submission_survive_new_service_instance(self):
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив', 'Нет', '3', 'Помогать',
                     '5 часов', 'Книги', 'Аккуратность', 'VK', 'Каталогизировал', 'Медиа']:
            await self.process(text)
            self.s = Service(self.db, self.vk, self.settings)
        async with self.db.pool.connection() as c:
            row = await (await c.execute('SELECT a.* FROM applications a JOIN users u ON u.id=a.user_id WHERE u.vk_id=%s', (self.vk_id,))).fetchone()
            count = await (await c.execute('SELECT count(*) n FROM answers WHERE application_id=%s', (row['id'],))).fetchone()
        self.assertEqual(row['status'], 'submitted')
        self.assertEqual(count['n'], 12)
        self.assertIsNotNone(row['consent_at'])

    async def test_slash_start_resumes_questionnaire(self):
        await self.process('/start')
        async with self.db.pool.connection() as c:
            app = await (await c.execute('SELECT state FROM applications WHERE group_id=%s', (self.group,))).fetchone()
        self.assertIsNotNone(app)
        self.assertEqual(app['state']['step'], 'consent')

    async def test_no_consent_text_starts_with_name(self):
        self.settings.consent_text = ''
        await self.process('Старт')
        async with self.db.pool.connection() as c:
            app = await (await c.execute('SELECT state FROM applications WHERE group_id=%s', (self.group,))).fetchone()
        self.assertEqual(app['state']['step'], 'name')

    async def test_only_staff_can_retry_failed_job(self):
        async with self.db.pool.connection() as c:
            jid = await self.db.enqueue(c, 'score', {'group_id': self.group}, str(uuid.uuid4()))
            await c.execute("UPDATE jobs SET status='failed',attempts=5 WHERE id=%s", (jid,))
        await self.process('/retry ' + str(jid))
        async with self.db.pool.connection() as c:
            status = await (await c.execute('SELECT status FROM jobs WHERE id=%s', (jid,))).fetchone()
            await c.execute("UPDATE users SET role='admin' WHERE vk_id=%s", (self.vk_id,))
        self.assertEqual(status['status'], 'failed')
        await self.process('/retry ' + str(jid))
        async with self.db.pool.connection() as c:
            status = await (await c.execute('SELECT status,attempts FROM jobs WHERE id=%s', (jid,))).fetchone()
        self.assertEqual(status, {'status': 'pending', 'attempts': 0})

    async def test_completed_application_has_contextual_question_snapshot(self):
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив и документы', 'Нет', '3', 'Помогать',
                     '5 часов', 'Книги', 'Аккуратность', 'VK', 'Проверял дважды', 'Медиа']:
            await self.process(text)
        async with self.db.pool.connection() as c:
            answer = await (await c.execute("SELECT question_text FROM answers an JOIN applications a ON a.id=an.application_id WHERE a.group_id=%s AND question_code='followup_ai_1'", (self.group,))).fetchone()
        self.assertIn('проверяли точность', answer['question_text'])

    async def test_resume_document_is_linked_without_advancing_empty_answer(self):
        await self.process('Старт')
        await self.process('Согласен')
        await self.process('', [{'type': 'doc', 'doc': {'id': 7, 'owner_id': self.vk_id, 'title': 'resume.pdf', 'size': 1234, 'ext': 'pdf'}}])
        async with self.db.pool.connection() as c:
            file = await (await c.execute("SELECT f.* FROM files f JOIN applications a ON a.id=f.application_id WHERE a.group_id=%s", (self.group,))).fetchone()
            app = await (await c.execute('SELECT state FROM applications WHERE group_id=%s', (self.group,))).fetchone()
        self.assertIsNotNone(file)
        self.assertEqual(file['kind'], 'resume')
        self.assertEqual(app['state']['step'], 'name')

    async def test_optional_consent_mode_does_not_save_first_message_attachment(self):
        await self.process('Старт', [{'type': 'doc', 'doc': {'id': 7, 'owner_id': self.vk_id, 'title': 'resume.pdf'}}])
        async with self.db.pool.connection() as c:
            count = await (await c.execute('SELECT count(*) n FROM files')).fetchone()
        self.assertEqual(count['n'], 0)

    async def test_question_snapshot_survives_new_question_version(self):
        from gim.dialog import QUESTIONS
        await self.process('Старт')
        await self.process('Согласен')
        await self.process('Иван')
        with patch.dict(QUESTIONS, {'experience': 'Изменённый вопрос'}):
            job = await self.process('История')
        async with self.db.pool.connection() as c:
            reply = await (await c.execute('SELECT payload FROM jobs WHERE dedupe_key=%s', (str(job['id']) + ':reply:0',))).fetchone()
        self.assertTrue(reply['payload']['text'].startswith('Расскажите об опыте'))

    async def test_export_uses_campaign_from_request_after_restart(self):
        import io
        from openpyxl import load_workbook
        from pathlib import Path
        await self.process('Старт')
        await self.process('Согласен')
        await self.process('Кандидат первого набора')
        async with self.db.pool.connection() as c:
            user = await (await c.execute("UPDATE users SET role='curator' WHERE vk_id=%s RETURNING *", (self.vk_id,))).fetchone()
            await self.db.enqueue(c, 'export', {'peer_id': self.vk_id, 'campaign': 'test'}, str(uuid.uuid4()), recipient_id=user['id'])
        job = await self.db.claim(kinds=['export'])
        self.settings.campaign = 'new-campaign'
        await self.s.export(job)
        async with self.db.pool.connection() as c:
            file = await (await c.execute('SELECT * FROM files WHERE id=%s', (job['id'],))).fetchone()
        workbook = load_workbook(io.BytesIO((Path(self.settings.storage_dir) / file['storage_key']).read_bytes()), read_only=True)
        rows = list(workbook.active.values)
        workbook.close()
        self.assertEqual(file['campaign'], 'test')
        self.assertEqual(len(rows), 2)

    async def test_export_can_be_opened_and_contains_saved_answers(self):
        import io
        from openpyxl import load_workbook
        from pathlib import Path
        await self.process('Старт')
        await self.process('Согласен')
        await self.process('Тестовый кандидат')
        async with self.db.pool.connection() as c:
            user = await (await c.execute("UPDATE users SET role='curator' WHERE vk_id=%s RETURNING *", (self.vk_id,))).fetchone()
            await self.db.enqueue(c, 'export', {'peer_id': self.vk_id}, str(uuid.uuid4()), recipient_id=user['id'])
        job = await self.db.claim(kinds=['export'])
        await self.s.export(job)
        async with self.db.pool.connection() as c:
            file = await (await c.execute('SELECT * FROM files WHERE id=%s', (job['id'],))).fetchone()
        data = (Path(self.settings.storage_dir) / file['storage_key']).read_bytes()
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        rows = list(workbook.active.values)
        workbook.close()
        self.assertEqual(rows[1][1], 'Тестовый кандидат')
        self.assertIn('Тестовый кандидат', rows[1][9])

    async def test_export_retry_uses_published_snapshot(self):
        async with self.db.pool.connection() as c:
            user = await (await c.execute("INSERT INTO users(vk_id,role) VALUES(%s,'curator') RETURNING *", (self.vk_id,))).fetchone()
            await self.db.enqueue(c, 'export', {'peer_id': self.vk_id}, str(uuid.uuid4()), recipient_id=user['id'])
        job = await self.db.claim(kinds=['export'])
        await self.s.export(job)
        async with self.db.pool.connection() as c:
            first = await (await c.execute('SELECT * FROM files WHERE id=%s', (job['id'],))).fetchone()
            await c.execute("UPDATE jobs SET lease_until=now()-interval '1 second' WHERE id=%s", (job['id'],))
        retry = await self.db.claim(kinds=['export'])
        await self.s.export(retry)
        async with self.db.pool.connection() as c:
            last = await (await c.execute('SELECT * FROM files WHERE id=%s', (job['id'],))).fetchone()
            sends = await (await c.execute("SELECT payload FROM jobs WHERE kind='send_file'")).fetchall()
        self.assertEqual(first['storage_key'], last['storage_key'])
        self.assertEqual(len(sends), 1)
        self.assertEqual(sends[0]['payload']['storage_key'], last['storage_key'])

    async def test_staff_reply_carries_permission_guard(self):
        async with self.db.pool.connection() as c:
            await c.execute("INSERT INTO users(vk_id,role) VALUES(%s,'curator')", (self.vk_id,))
        job = await self.process('/top')
        async with self.db.pool.connection() as c:
            reply = await (await c.execute("SELECT * FROM jobs WHERE dedupe_key=%s", (str(job['id']) + ':reply:0',))).fetchone()
        self.assertIsNotNone(reply['recipient_id'])

    async def test_revoked_staff_cannot_receive_export(self):
        async with self.db.pool.connection() as c:
            user = await (await c.execute("INSERT INTO users(vk_id,role) VALUES(%s,'curator') RETURNING *", (self.vk_id,))).fetchone()
            jid = await self.db.enqueue(c, 'send', {'peer_id': self.vk_id, 'text': 'private'}, str(uuid.uuid4()), recipient_id=user['id'])
            await c.execute('UPDATE users SET active=false WHERE id=%s', (user['id'],))
            job = await (await c.execute('SELECT * FROM jobs WHERE id=%s', (jid,))).fetchone()
        await self.s.handle(job)
        self.vk.send.assert_not_awaited()

    async def test_stale_export_cannot_publish_file(self):
        async with self.db.pool.connection() as c:
            user = await (await c.execute("INSERT INTO users(vk_id,role) VALUES(%s,'curator') RETURNING *", (self.vk_id,))).fetchone()
            jid = await self.db.enqueue(c, 'export', {'peer_id': self.vk_id}, str(uuid.uuid4()), recipient_id=user['id'])
            job = await (await c.execute('SELECT * FROM jobs WHERE id=%s', (jid,))).fetchone()
        # Pending job has never been leased: simulates a handler with a stale token.
        job['lease_token'] = uuid.uuid4()
        await self.s.export(job)
        async with self.db.pool.connection() as c:
            file = await (await c.execute('SELECT id FROM files WHERE id=%s', (jid,))).fetchone()
        self.assertIsNone(file)

    async def test_scoring_contract_stores_evidence_and_omits_future_outcomes(self):
        import httpx
        import json
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив', 'Нет', '3', 'Помогать',
                     '5 часов', 'Книги', 'Аккуратность', 'VK', 'Каталогизировал', 'Медиа']:
            await self.process(text)
        job = await self.db.claim(kinds=['score'])
        self.settings.scoring_url = 'https://scoring.test/evaluate'
        async def endpoint(request):
            payload = json.loads(request.content)
            self.assertNotIn('historical_outcomes', payload)
            self.assertNotIn('name', {a['question_code'] for a in payload['answers']})
            evidence = next(a for a in payload['answers'] if a['question_code'] == 'experience')
            return httpx.Response(200, json={'score': 2, 'rubric_version': 'team-v1',
                'rubric_snapshot': {'criteria': [{'code': 'experience'}]},
                'criteria_results': [{'criterion': 'experience', 'points': 2, 'reason': 'Опыт архива', 'answer_id': evidence['id'], 'quote': 'Архив'}],
                'summary': 'Релевантный опыт', 'interview_questions': ['Какие документы обрабатывали?'], 'recommended_directions': ['funds']})
        factory = httpx.AsyncClient
        with patch('gim.service.httpx.AsyncClient', lambda **kwargs: factory(transport=httpx.MockTransport(endpoint), **kwargs)):
            await self.s.score(job)
        async with self.db.pool.connection() as c:
            evaluation = await (await c.execute('SELECT * FROM evaluations WHERE application_id=%s', (job['application_id'],))).fetchone()
        self.assertEqual(evaluation['score'], 2)
        self.assertEqual(evaluation['criteria_results'][0]['quote'], 'Архив')
        async with self.db.pool.connection() as c:
            await c.execute("UPDATE users SET role='curator' WHERE vk_id=%s", (self.vk_id,))
        await self.process('/score ' + str(job['application_id']) + ' 1.5 Требуется собеседование')
        async with self.db.pool.connection() as c:
            row = (await self.s.rows(c))[0]
        self.assertEqual(row['score'], 1.5)
        self.assertEqual(row['source'], 'human')
        self.assertEqual(row['criteria_results'][0]['quote'], 'Архив')
        self.assertEqual(row['summary'], 'Релевантный опыт')
        self.assertEqual(row['manual_reason'], 'Требуется собеседование')

    async def test_followup_service_failure_preserves_answers_without_pattern_fallback(self):
        import httpx
        self.auto_ai = False
        self.settings.followup_url = 'https://questions.test/select'
        self.settings.service_token = ''
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив и документы', 'Нет', '3',
                     'Помогать', '5 часов', 'Книги', 'Аккуратность', 'VK']:
            await self.process(text)
        job = await self.db.claim(kinds=['prepare_followups'])
        factory = httpx.AsyncClient
        with patch('gim.service.httpx.AsyncClient', lambda **kwargs: factory(transport=httpx.MockTransport(
                lambda request: httpx.Response(503)), **kwargs)):
            with self.assertRaises(httpx.HTTPStatusError):
                await self.s.prepare_followups(job)
        async with self.db.pool.connection() as c:
            app = await (await c.execute('SELECT * FROM applications WHERE group_id=%s', (self.group,))).fetchone()
        self.assertEqual(app['state']['step'], 'planning')
        self.assertEqual(app['revision'], 10)
        self.assertNotIn('followup_ai_1', app['questionnaire'])

    async def test_limit_finishes_clarifications_without_another_model_call(self):
        self.settings.max_followups = 1
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив', 'Нет', '3',
                     'Помогать', '5 часов', 'Книги', 'Аккуратность', 'VK', 'Проверял дважды']:
            await self.process(text)
        async with self.db.pool.connection() as c:
            app = await (await c.execute('SELECT state FROM applications WHERE group_id=%s', (self.group,))).fetchone()
            count = await (await c.execute("SELECT count(*) n FROM jobs WHERE kind='prepare_followups'")).fetchone()
        self.assertEqual(app['state']['step'], 'expectations')
        self.assertEqual(count['n'], 1)

    async def test_extra_message_while_model_thinks_is_saved_and_stale_result_ignored(self):
        self.auto_ai = False
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив', 'Нет', '3',
                     'Помогать', '5 часов', 'Книги', 'Аккуратность', 'VK']:
            await self.process(text)
        old_job = await self.db.claim(kinds=['prepare_followups'])
        # A concurrent incoming event is processed while HTTP generation is in flight.
        await self.process('Дополню: ещё составлял описи монет')
        await self.s.prepare_followups(old_job)
        async with self.db.pool.connection() as c:
            app = await (await c.execute('SELECT * FROM applications WHERE group_id=%s', (self.group,))).fetchone()
            extra = await (await c.execute("SELECT answer_text FROM answers WHERE question_kind='supplement'")).fetchone()
        self.assertEqual(app['revision'], 11)
        self.assertEqual(app['state']['step'], 'planning')
        self.assertIn('описи монет', extra['answer_text'])

    async def test_queued_addition_prevents_publication_of_stale_question(self):
        import httpx
        self.auto_ai = False
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив', 'Нет', '3',
                     'Помогать', '5 часов', 'Книги', 'Аккуратность', 'VK']:
            await self.process(text)
        ai_job = await self.db.claim(kinds=['prepare_followups'])
        async def endpoint(request):
            await self.db.ingest({'type': 'message_new', 'group_id': self.group, 'event_id': str(uuid.uuid4()),
                'object': {'message': {'from_id': self.vk_id, 'peer_id': self.vk_id, 'text': 'Дополнение ещё до вопроса'}}})
            return httpx.Response(200, json={'action': 'ask', 'direction': 'funds', 'question': 'Какие предметы вы описывали?'})
        factory = httpx.AsyncClient
        with patch('gim.service.httpx.AsyncClient', lambda **kwargs: factory(transport=httpx.MockTransport(endpoint), **kwargs)):
            with self.assertRaises(RuntimeError):
                await self.s.prepare_followups(ai_job)
        async with self.db.pool.connection() as c:
            app = await (await c.execute('SELECT state FROM applications WHERE group_id=%s', (self.group,))).fetchone()
        self.assertEqual(app['state']['step'], 'planning')

    async def test_scoring_adapter_cannot_invent_criteria_or_skip_evidence(self):
        import httpx
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив', 'Нет', '3',
                     'Помогать', '5 часов', 'Книги', 'Аккуратность', 'VK', 'Каталогизировал', 'Медиа']:
            await self.process(text)
        ai_job = await self.db.claim(kinds=['score'])
        self.settings.scoring_url = 'https://scoring.test/evaluate'
        result = {'score': 3, 'rubric_version': 'invented', 'rubric_snapshot': {'anything': True},
                  'criteria_results': [{'criterion': 'invented', 'reason': 'Потому что так'}],
                  'summary': 'Без доказательств', 'interview_questions': [], 'recommended_directions': ['funds']}
        factory = httpx.AsyncClient
        with patch('gim.service.httpx.AsyncClient', lambda **kwargs: factory(transport=httpx.MockTransport(
                lambda request: httpx.Response(200, json=result)), **kwargs)):
            with self.assertRaises(ValueError):
                await self.s.score(ai_job)

    async def test_direct_model_builds_final_summary_from_full_dialog(self):
        import httpx
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив', 'Нет', '3',
                     'Помогать', '5 часов', 'Книги', 'Аккуратность', 'VK', 'Каталогизировал', 'Медиа']:
            await self.process(text)
        ai_job = await self.db.claim(kinds=['score'])
        self.settings.llm_api_key, self.settings.llm_model = 'test-secret', 'test-model'
        self.settings.llm_base_url = 'https://model.test/v1'
        async def endpoint(request):
            payload = json.loads(json.loads(request.content)['messages'][1]['content'])
            self.assertEqual(payload['answers'][-1]['question_code'], 'expectations')
            evidence = next(a for a in payload['answers'] if a['question_code'] == 'experience')
            result = {'score': 2, 'criteria_results': [{'criterion': 'experience', 'points': 2,
                'reason': 'Работал в архиве', 'answer_id': evidence['id'], 'quote': 'Архив'}],
                'summary': 'Итог после всех ответов', 'interview_questions': ['Какие документы?'], 'recommended_directions': ['funds']}
            return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(result)}}]})
        factory = httpx.AsyncClient
        with patch('gim.llm.httpx.AsyncClient', lambda **kwargs: factory(transport=httpx.MockTransport(endpoint), **kwargs)):
            await self.s.score(ai_job)
        async with self.db.pool.connection() as c:
            evaluation = await (await c.execute('SELECT * FROM evaluations WHERE application_id=%s', (ai_job['application_id'],))).fetchone()
            frozen = await (await c.execute('SELECT payload FROM jobs WHERE id=%s', (ai_job['id'],))).fetchone()
        self.assertEqual(evaluation['summary'], 'Итог после всех ответов')
        self.assertEqual(frozen['payload']['rubric']['version'], 'team-v1')

    async def test_pre_upgrade_job_without_revision_remains_processable(self):
        import httpx
        self.auto_ai = False
        for text in ['Старт', 'Согласен', 'Иван', 'История', 'Архив', 'Нет', '3',
                     'Помогать', '5 часов', 'Книги', 'Аккуратность', 'VK']:
            await self.process(text)
        ai_job = await self.db.claim(kinds=['prepare_followups'])
        ai_job['payload'].pop('revision')
        async with self.db.pool.connection() as c:
            await c.execute("UPDATE applications SET state=state-'max_followups' WHERE id=%s", (ai_job['application_id'],))
        factory = httpx.AsyncClient
        with patch('gim.service.httpx.AsyncClient', lambda **kwargs: factory(transport=httpx.MockTransport(
                lambda request: httpx.Response(200, json={'action': 'finish'})), **kwargs)):
            await self.s.prepare_followups(ai_job)
        async with self.db.pool.connection() as c:
            app = await (await c.execute('SELECT state FROM applications WHERE group_id=%s', (self.group,))).fetchone()
        self.assertEqual(app['state']['step'], 'expectations')
